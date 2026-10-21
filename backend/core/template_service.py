"""Checklist schema validation, revisioning, and custom-template report export."""
import json
import re
from copy import deepcopy
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.files.storage import default_storage
from django.db import transaction
from django.db.models import Max

FIELD_TYPES = {'text', 'textarea', 'number', 'date', 'select', 'inspection', 'table', 'photo'}
IDENTIFIER = re.compile(r'^[a-zA-Z][a-zA-Z0-9_-]{0,79}$')


def default_definition():
    return json.loads(Path(__file__).with_name('checklist_defaults.json').read_text())


def validate_definition(definition):
    if not isinstance(definition, dict) or not isinstance(definition.get('sections'), list):
        raise ValidationError('A checklist must contain a sections list.')
    if not 1 <= len(definition['sections']) <= 100:
        raise ValidationError('Provide between 1 and 100 sections.')
    if 'standard_catalog' in definition:
        from .standard_checklist_service import standard_metadata
        standard_metadata(definition)
    if 'wizard_context' in definition:
        from .standard_checklist_service import wizard_metadata
        wizard_metadata(definition)
    section_ids, field_ids = set(), set()
    for section in definition['sections']:
        if not isinstance(section, dict):
            raise ValidationError('Each section must be an object.')
        sid = section.get('id', '')
        if not isinstance(sid, str) or not IDENTIFIER.fullmatch(sid) or sid in section_ids:
            raise ValidationError('Section identifiers must be valid and unique.')
        section_ids.add(sid)
        if not isinstance(section.get('title'), str) or not section['title'].strip() or len(section['title']) > 160:
            raise ValidationError('Every section needs a title of at most 160 characters.')
        if not isinstance(section.get('group', ''), str) or len(section.get('group', '')) > 160:
            raise ValidationError('A section group must be a short text label.')
        fields = section.get('fields')
        if not isinstance(fields, list) or not 1 <= len(fields) <= 300:
            raise ValidationError('Every section must contain between 1 and 300 fields.')
        for field in fields:
            if not isinstance(field, dict):
                raise ValidationError('Each field must be an object.')
            fid = field.get('id', '')
            if not isinstance(fid, str) or not IDENTIFIER.fullmatch(fid) or fid in field_ids:
                raise ValidationError('Field identifiers must be valid and unique across the checklist.')
            field_ids.add(fid)
            if not isinstance(field.get('label'), str) or not field['label'].strip() or len(field['label']) > 3000:
                raise ValidationError('Every field needs a label of at most 3000 characters.')
            if field.get('type') not in FIELD_TYPES:
                raise ValidationError('Unknown field type.')
            if field.get('system') and (field['system'] != 'site_id' or field['type'] != 'text'):
                raise ValidationError('The linked Site ID field must remain a text field.')
            if not isinstance(field.get('required', False), bool):
                raise ValidationError('Required must be true or false.')
            if 'evidence_rules' in field:
                from .evidence_rules import validate_evidence_rules
                try:validate_evidence_rules(field['evidence_rules'])
                except ValueError as error:raise ValidationError(str(error)) from error
            if field.get('type') == 'select':
                options = field.get('options')
                if not isinstance(options, list) or not 1 <= len(options) <= 100 or any(not isinstance(x, str) or not x.strip() or len(x)>200 for x in options):
                    raise ValidationError('Dropdown fields require a nonempty list of choices.')
            if field.get('type') == 'table':
                columns = field.get('columns')
                if not isinstance(columns, list) or not 1 <= len(columns) <= 20:
                    raise ValidationError('Repeating tables require between 1 and 20 columns.')
                ids = set()
                for column in columns:
                    if not isinstance(column, dict) or not isinstance(column.get('id'), str) or not IDENTIFIER.fullmatch(column['id']) or column['id'] in ids or not isinstance(column.get('label'), str) or not column['label'].strip():
                        raise ValidationError('Table columns need unique identifiers and labels.')
                    ids.add(column['id'])
    if len(field_ids) > 3000:
        raise ValidationError('A checklist can contain at most 3000 fields.')


def field_map(checklist):
    return {f['id']: f for s in checklist.schema_snapshot.get('sections', []) for f in s['fields']}


def validate_answer(field, value):
    kind = field['type']
    if kind == 'table':
        if not isinstance(value, list) or len(value) > 500:
            raise ValidationError('Repeating table data must be a list of at most 500 rows.')
        columns = {c['id'] for c in field['columns']}
        for row in value:
            if not isinstance(row, dict) or set(row) - columns or any(not isinstance(v, str) or len(v)>5000 for v in row.values()):
                raise ValidationError('Invalid repeating-table row.')
        return
    if kind == 'photo':
        raise ValidationError('Use the photo upload action for evidence files.')
    if kind == 'inspection':
        if not isinstance(value, dict) or value.get('result', '') not in ['', 'Checked', 'Defect found', 'Not applicable'] or not isinstance(value.get('remarks', ''), str) or len(value.get('remarks', '')) > 10000:
            raise ValidationError('Invalid inspection result.')
        if value.get('result') == 'Not applicable' and not value.get('remarks', '').strip():
            raise ValidationError('Provide a reason for Not applicable.')
        return
    if not isinstance(value, str) or len(value) > 10000:
        raise ValidationError('Field values must be text of at most 10000 characters.')
    if value and kind == 'number':
        try:
            if not Decimal(value).is_finite():
                raise InvalidOperation()
        except InvalidOperation:
            raise ValidationError('Enter a valid finite number.')
    if value and kind == 'date':
        try:
            date.fromisoformat(value)
        except ValueError:
            raise ValidationError('Enter a date in YYYY-MM-DD format.')
    if value and kind == 'select' and value not in field.get('options', []):
        raise ValidationError('Select a configured dropdown value.')



def declared_evidence_missing(checklist,field,links,record_cache):
    """Read saved results only; inference is forbidden in validation/read requests."""
    from .evidence_rules import validate_evidence_rules
    from .field_photo_service import schema_digest
    from .evidence_learning import digest
    from .evidence_models import EvidenceAnalysisRecord
    try:rules=validate_evidence_rules(field.get('evidence_rules'))
    except ValueError:return [field['label']+' invalid evidence requirements']
    active={key:value for key,value in rules.items() if key not in {'version','notes'} and value not in (False,None,[],{},'any')}
    if not active:return []
    minimum=rules.get('min_photos',1)
    photos={link.photo_id:link.photo for link in links}
    if len(photos)<minimum:return [field['label']+' requires '+str(minimum)+' distinct confirmed source photos']
    if set(active)<={'min_photos'}:return []
    eligible=0;contradictory=False
    for photo_id,photo in photos.items():
        if photo_id not in record_cache:
            record_cache[photo_id]=EvidenceAnalysisRecord.objects.filter(checklist_id=checklist.pk,project_id=checklist.project_id,company_id=checklist.project.company_id,site_id=checklist.site_id,photo_id=photo_id,photo_sha256=photo.sha256,schema_digest=schema_digest(checklist)).order_by('-pk').first()
        record=record_cache[photo_id]
        if not record:continue
        try:
            if digest(record.payload)!=record.payload_digest:continue
        except ValidationError:continue
        assessments=record.payload.get('assessments',[]) if isinstance(record.payload,dict) else []
        assessment=next((item for item in assessments if isinstance(item,dict) and item.get('field_id')==field['id']),None)
        if not assessment:continue
        if assessment.get('outcome')=='contradictory':contradictory=True
        if assessment.get('outcome')=='supported' and not assessment.get('hard_blocked'):eligible+=1
    result=[]
    if contradictory:result.append(field['label']+' has contradictory saved source evidence')
    if eligible<minimum:result.append(field['label']+' evidence requirements need '+str(minimum)+' supported confirmed source photos; '+str(eligible)+' available')
    return result

def missing_required(checklist, evidence_checklist=None):
    from .field_photo_service import confirmed_placements
    from .models import PhotoJob
    source=evidence_checklist or checklist
    linked = confirmed_placements(evidence_checklist or checklist)
    answers = (checklist.answer_data or {}).get('custom', {})
    missing = []
    if getattr(source, 'checklist_template_id', None):
        from .survey_processing import require_engineer_review
        try:
            require_engineer_review(source)
        except ValidationError as exc:
            missing.extend(exc.messages)
    if getattr(source,'pk',None) and PhotoJob.objects.filter(checklist_id=source.pk,
            schema=source.schema_snapshot,status__in=['queued','running']).exists():
        missing.append('Photo analysis is still active; wait for processing to finish before submission or approval')
    record_cache={}
    for fid, field in field_map(checklist).items():
        if field.get('evidence_rules'):
            missing.extend(declared_evidence_missing(evidence_checklist or checklist,field,linked.get(fid,[]),record_cache))
        if field.get('photo_required') and not ((checklist.image_data or {}).get(fid) or linked.get(fid)):
            missing.append(field['label'] + ' photo evidence')
        if not field.get('required'):
            continue
        value = checklist.site_id if field.get('system') == 'site_id' else answers.get(fid)
        if field['type'] == 'photo':
            value = (checklist.image_data or {}).get(fid) or linked.get(fid)
        if field['type'] == 'inspection':
            value = value.get('result') if isinstance(value, dict) else None
        if field['type'] == 'table' and isinstance(value, list):
            value = any(any(str(cell).strip() for cell in row.values()) for row in value)
        if not value or (isinstance(value, str) and not value.strip()):
            missing.append(field['label'])
    return missing


@transaction.atomic
def duplicate_revision(template, user):
    from .models import ChecklistTemplate
    # Lock the stable first revision to serialize revision-number allocation.
    list(ChecklistTemplate.objects.select_for_update().filter(family_id=template.family_id).order_by('version')[:1])
    next_version = (ChecklistTemplate.objects.filter(family_id=template.family_id).aggregate(n=Max('version'))['n'] or 0) + 1
    return ChecklistTemplate.objects.create(family_id=template.family_id, version=next_version, name=template.name, definition=deepcopy(template.definition), checklist_required=template.checklist_required, company_logo=template.company_logo.name, client_left_name=template.client_left_name, client_right_name=template.client_right_name, client_left_logo=template.client_left_logo.name, client_right_logo=template.client_right_logo.name, created_by=user)


def report_value(value):
    if isinstance(value, dict):
        return ' | '.join(f'{k}: {v}' for k, v in value.items())
    return str(value or '')


def safe_excel_text(value):
    value = report_value(value)
    return "'"+value if value.startswith(('=', '+', '-', '@')) else value


def export_custom_excel(checklist):
    """Build a report from the frozen schema, current answers, and client branding."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.drawing.image import Image
    from .field_photo_service import confirmed_placements
    linked = confirmed_placements(checklist)
    workbook = Workbook()
    ws = workbook.active
    ws.title = 'Site survey'
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 66
    ws.column_dimensions['C'].width = 55
    ws.column_dimensions['D'].width = 50
    logos = checklist.branding_snapshot.get('clients', [])
    for i, client in enumerate(logos):
        target = 'A1' if i == 0 else 'D1'
        ws[target] = safe_excel_text(client.get('name', ''))
        path = client.get('logo')
        if path and default_storage.exists(path):
            with default_storage.open(path, 'rb') as fh:
                im = Image(BytesIO(fh.read()))
            ratio = min(140 / im.width, 48 / im.height)
            im.width, im.height = im.width * ratio, im.height * ratio
            ws.add_image(im, 'A2' if i == 0 else 'D2')
    ws.row_dimensions[2].height = 40
    ws.append([])
    ws['B4'] = safe_excel_text(checklist.schema_snapshot.get('title', 'Site survey'))
    ws['B4'].font = Font(size=17, bold=True)
    ws['B5'], ws['C5'] = 'Site ID', safe_excel_text(checklist.site_id)
    ws['B6'], ws['C6'] = 'Template revision', str(checklist.checklist_template)
    ws['B7'], ws['C7'] = 'Status', checklist.get_status_display()
    ws['B8'], ws['C8'] = 'Review comment', safe_excel_text(checklist.comment)
    ws['B9'], ws['C9'] = 'Reviewed by', safe_excel_text(checklist.comment_by.get_username() if checklist.comment_by else '')
    answers = (checklist.answer_data or {}).get('custom', {})
    row = 11
    for section in checklist.schema_snapshot['sections']:
        ws.cell(row, 2, safe_excel_text(section['title']))
        for c in range(1, 5):
            ws.cell(row, c).fill = PatternFill('solid', fgColor='EAF2DF')
        ws.cell(row, 2).font = Font(bold=True)
        row += 1
        for field in section['fields']:
            value = checklist.site_id if field.get('system') == 'site_id' else answers.get(field['id'], '')
            ws.cell(row, 1, 'Required' if field.get('required') else '')
            ws.cell(row, 2, safe_excel_text(field['label']))
            if field['type'] == 'table' and isinstance(value, list):
                ws.cell(row, 3, safe_excel_text(' | '.join(c['label'] for c in field['columns'])))
                for item in value:
                    row += 1
                    ws.cell(row, 3, safe_excel_text(' | '.join(item.get(c['id'], '') for c in field['columns'])))
            else:
                ws.cell(row, 3, safe_excel_text(value))
            paths = (checklist.image_data or {}).get(field['id'], [])
            for path in paths:
                if not default_storage.exists(path):
                    continue
                row += 1
                with default_storage.open(path, 'rb') as fh:
                    im = Image(BytesIO(fh.read()))
                scale = min(240 / im.width, 180 / im.height)
                im.width, im.height = im.width * scale, im.height * scale
                ws.add_image(im, f'D{row}')
                ws.row_dimensions[row].height = max(im.height * .75, 30)
            for placement in linked.get(field['id'], []):
                photo=placement.photo
                if not photo.preview or not photo.preview.storage.exists(photo.preview.name):
                    continue
                row += 1
                ws.cell(row, 3, safe_excel_text(f'Photo {photo.pk}: {photo.filename} ({placement.status})'))
                with photo.preview.open('rb') as fh:
                    im=Image(BytesIO(fh.read()))
                scale=min(240/im.width,180/im.height)
                im.width,im.height=im.width*scale,im.height*scale
                ws.add_image(im,f'D{row}')
                ws.row_dimensions[row].height=max(im.height*.75,30)
            row += 1
    for cells in ws:
        for cell in cells:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'C11'
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    review = workbook.create_sheet('Section approvals')
    review.append(['Section', 'Status', 'Reviewed by', 'Date'])
    for section in checklist.schema_snapshot['sections']:
        approval = checklist.section_approvals.get(section['id'], {})
        review.append([safe_excel_text(section['title']), 'Approved' if checklist.status == checklist.Status.FINAL or approval else 'Pending', safe_excel_text(approval.get('name', '')), approval.get('time', '')])
    history = workbook.create_sheet('Review history')
    history.append(['Date', 'Reviewer', 'Decision', 'Section IDs', 'Comment'])
    for event in checklist.review_history:
        history.append([event.get('time', ''), safe_excel_text(event.get('name', '')), event.get('action', ''), safe_excel_text(', '.join(event.get('sections', []))), safe_excel_text(event.get('comment', ''))])
    if linked:
        evidence=workbook.create_sheet('Evidence locations')
        evidence.append(['Field ID','Section ID','Photo ID','Filename','Group','Placement status','Source SHA256','Reviewed by','Reviewed at'])
        for fid,items in linked.items():
            for item in items:
                evidence.append([safe_excel_text(fid),safe_excel_text(item.section_id),item.photo_id,safe_excel_text(item.photo.filename),safe_excel_text(item.photo.group),item.status,item.photo.sha256,item.reviewed_by_id,item.reviewed_at.isoformat() if item.reviewed_at else ''])
        evidence.freeze_panes='A2'
        for column in 'ABCDEFGHI':
            evidence.column_dimensions[column].width=25
        for cell in evidence[1]:
            cell.font=Font(bold=True)
    for sheet in [review, history]:
        sheet.freeze_panes = 'A2'
        for column in ['A', 'B', 'C', 'D', 'E']:
            sheet.column_dimensions[column].width = 30 if column != 'E' else 70
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
