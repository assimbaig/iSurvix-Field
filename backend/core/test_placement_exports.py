"""Confirmed field evidence is visible, exportable and required-photo aware."""
import hashlib,tempfile
from pathlib import Path
from io import BytesIO
from django.contrib.auth.models import User
from django.core.files.base import ContentFile
from django.test import TestCase,override_settings,RequestFactory
from django.urls import reverse
from PIL import Image
from openpyxl import load_workbook
from .models import Project,Profile,Checklist,ChecklistTemplate,PhotoEvidence,PhotoPlacement
from .field_photo_service import schema_digest
from .template_service import missing_required,export_custom_excel
from .template_views import survey_context

@override_settings(ROOT_URLCONF='checklist.urls_local',ALLOWED_HOSTS=['testserver'],SECURE_SSL_REDIRECT=False,PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class PlacementExportTests(TestCase):
    def setUp(self):
        folder=Path(__file__).resolve().parents[2]/'.local/tmp';folder.mkdir(parents=True,exist_ok=True)
        self.media=tempfile.TemporaryDirectory(dir=folder)
        self.override=override_settings(MEDIA_ROOT=self.media.name);self.override.enable()
        self.addCleanup(self.media.cleanup);self.addCleanup(self.override.disable)
        schema={'title':'Synthetic mapped evidence','sections':[{'id':'civil','title':'Civil','fields':[{'id':'overview','label':'Civil overview','type':'photo','required':True}]},{'id':'electrical','title':'Electrical','fields':[{'id':'notes','label':'Recorded reading','type':'text','required':True}]}]}
        template=ChecklistTemplate.objects.create(name='Mapped evidence test',definition=schema,published=True)
        self.project=Project.objects.create(name='Mapped evidence project',checklist_template=template)
        self.engineer=self.person('mapping_engineer',Profile.Roles.ENGINEER)
        self.lead=self.person('mapping_lead',Profile.Roles.TEAM_LEAD)
        self.quality=self.person('mapping_quality',Profile.Roles.QUALITY_ENGINEER)
        self.survey=Checklist.objects.create(user=self.engineer,project=self.project,site_id='SYNTHETIC-MAPPED')
        raw=BytesIO();Image.new('RGB',(32,32),'blue').save(raw,format='PNG');data=raw.getvalue()
        self.photo=PhotoEvidence(checklist=self.survey,uploaded_by=self.engineer,filename='=synthetic.png',sha256=hashlib.sha256(data).hexdigest(),group='Civil')
        self.photo.original.save('synthetic.bin',ContentFile(data),save=False)
        self.photo.preview.save('synthetic.png',ContentFile(data),save=False);self.photo.save()
        self.placement=PhotoPlacement.objects.create(photo=self.photo,field_id='overview',section_id='civil',status='proposed',schema_digest=schema_digest(self.survey),reason='Synthetic source proposal')
    def person(self,name,role):
        user=User.objects.create_user(name,password='Synthetic-mapping-894!')
        user.profile.role=role;user.profile.project=self.project;user.profile.save();return user
    def test_only_confirmed_locations_fulfil_requirements_and_export_images(self):
        self.assertIn('Civil overview',missing_required(self.survey))
        proposed=survey_context(self.survey,RequestFactory().get('/'))[0]['fields'][0]
        self.assertEqual(len(proposed['proposed_photos']),1);self.assertEqual(proposed['linked_photos'],[])
        self.placement.status='confirmed';self.placement.save()
        self.assertNotIn('Civil overview',missing_required(self.survey))
        field=survey_context(self.survey,RequestFactory().get('/'))[0]['fields'][0]
        self.assertEqual(len(field['linked_photos']),1);self.assertEqual(field['images'],[])
        workbook=load_workbook(BytesIO(export_custom_excel(self.survey)))
        self.assertEqual(len(workbook['Site survey']._images),1)
        self.assertEqual(workbook['Evidence locations']['D2'].value,"'=synthetic.png")
        self.survey.refresh_from_db();self.assertEqual(self.survey.image_data,{});self.assertEqual(self.survey.answer_data,{})
    def test_stale_mapping_cannot_satisfy_changed_checklist_or_export(self):
        self.placement.status='manual';self.placement.save()
        self.survey.schema_snapshot['sections'][0]['fields'][0]['label']='Revised required evidence'
        self.survey.save()
        self.assertIn('Revised required evidence',missing_required(self.survey))
        workbook=load_workbook(BytesIO(export_custom_excel(self.survey)))
        self.assertEqual(len(workbook['Site survey']._images),0)
        self.assertNotIn('Evidence locations',workbook.sheetnames)
    def test_partial_approval_uses_full_source_schema_for_confirmed_evidence(self):
        self.placement.status='confirmed';self.placement.save()
        self.survey.status=Checklist.Status.SUBMITTED;self.survey.save()
        self.survey.quality_reviewer=self.quality;self.survey.save()
        from .quality_workflow import review_quality
        review_quality(self.quality,self.survey,'approve',expected_updated_at=self.survey.updated_at.isoformat())
        self.client.force_login(self.lead)
        response=self.client.post(reverse('custom_checklist_review',args=[self.survey.pk]),{'action':'partial','sections':['civil']})
        self.assertEqual(response.status_code,302)
        self.survey.refresh_from_db();self.assertEqual(self.survey.status,Checklist.Status.PARTIAL)
        self.assertIn('civil',self.survey.section_approvals)
        self.assertNotIn('electrical',self.survey.section_approvals)
    def test_legacy_field_images_still_satisfy_required_photo(self):
        self.survey.image_data={'overview':[self.photo.preview.name]};self.survey.save()
        self.assertNotIn('Civil overview',missing_required(self.survey))
        workbook=load_workbook(BytesIO(export_custom_excel(self.survey)))
        self.assertEqual(len(workbook['Site survey']._images),1)
    def test_explicit_photo_requirement_is_distinct_from_inspection_answer(self):
        field=self.survey.schema_snapshot['sections'][0]['fields'][0]
        field.update(type='inspection',required=False,photo_required=True)
        self.survey.save()
        self.assertIn('Civil overview photo evidence',missing_required(self.survey))
        self.placement.schema_digest=schema_digest(self.survey)
        self.placement.status='manual';self.placement.save()
        self.assertNotIn('Civil overview photo evidence',missing_required(self.survey))
