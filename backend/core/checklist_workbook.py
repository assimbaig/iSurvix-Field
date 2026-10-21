"""Strict, file-free XLSX interchange for unpublished checklist definitions.

Workbook cells are data only: formulas, macros, external relationships, unknown
conditions and operational records are rejected. Blank optional cells mean absent;
extra_json may preserve an explicitly empty descriptive property on export.
"""
from copy import deepcopy
from io import BytesIO
import hashlib
import json
import math
import re
import zipfile
from xml.etree import ElementTree
from django.core.exceptions import ValidationError
from .template_service import FIELD_TYPES, IDENTIFIER, validate_definition
from .evidence_rules import validate_evidence_rules
from .standard_checklists import CATALOG, ROOM_TYPES, STAGES, STAGE_LABELS, _rooms

MAX_BYTES=5*1024*1024
MAX_EXPANDED=50*1024*1024
RULE_COLUMNS=('version','required_subjects','any_subjects','required_text','forbidden_subjects','expected_site_id','require_readable_text','require_qr','require_gps','require_capture_time','min_photos','minimum_quality','notes')
HEADERS={
 'Guide':['topic','instructions'],
 'Checklists':['checklist_key','name','industry','workflow_type','description','standard_key','stages','extra_json'],
 'Sections':['checklist_key','section_id','title','group','description','purpose','extra_json'],
 'Items':['checklist_key','section_id','field_id','label','type','required','photo_required','system','description','room_type','room_name','floor','extra_json'],
 'Choices':['checklist_key','field_id','choice','order'],
 'TableColumns':['checklist_key','field_id','column_id','label','order','extra_json'],
 'EvidenceRules':['checklist_key','field_id',*RULE_COLUMNS],
 'Rooms':['checklist_key','type','name','floor'],
 '_Lists':['list','value'],
}
ROOT_EXTRA={'title','description','technical_notice','topic_sources','wizard_context'}
SECTION_EXTRA={'group','description','purpose','room'}
FIELD_EXTRA={'description','unit','equipment_type','operator','excel_row','group','purpose','room_context'}
COLUMN_EXTRA={'description','unit'}
WORKFLOWS={'telecom_survey':'telecom','telecom_inspection':'telecom','construction_inspection':'construction'}

class WorkbookValidationError(ValueError):
    def __init__(self,errors):
        self.errors=errors
        super().__init__('; '.join(f"{e['sheet']} row {e['row']} {e['column']}: {e['message']}" for e in errors[:10]))

def _error(sheet,row,column,message):
    raise WorkbookValidationError([{'sheet':sheet,'row':row,'column':column,'message':message}])

def _json(value,sheet,row,column,kind=dict):
    if value in (None,''):return kind()
    if not isinstance(value,str):_error(sheet,row,column,'Use JSON text.')
    def unique(pairs):
        result={}
        for key,item in pairs:
            if key in result:raise ValueError('Duplicate JSON property.')
            result[key]=item
        return result
    try:
        result=json.loads(value,object_pairs_hook=unique,parse_constant=lambda value:(_ for _ in ()).throw(ValueError('Non-finite JSON number.')))
    except (ValueError,TypeError,RecursionError):_error(sheet,row,column,'Invalid JSON or duplicate property.')
    if not isinstance(result,kind):_error(sheet,row,column,f'Expected a JSON {kind.__name__}.')
    return result

def _text(value,sheet,row,column,limit=10000,required=False):
    if value is None:value=''
    if not isinstance(value,str) or len(value)>limit or any(ord(c)<32 and c not in '\n\r\t' for c in value):_error(sheet,row,column,'Expected bounded plain text.')
    if required and not value.strip():_error(sheet,row,column,'This value is required.')
    return value

def _id(value,sheet,row,column):
    value=_text(value,sheet,row,column,80,True)
    if not IDENTIFIER.fullmatch(value):_error(sheet,row,column,'Use a letter followed by letters, digits, underscores or hyphens (maximum 80).')
    return value

def _bool(value,sheet,row,column):
    if type(value) is bool:return value
    if isinstance(value,str) and value.upper() in {'TRUE','FALSE'}:return value.upper()=='TRUE'
    _error(sheet,row,column,'Use TRUE or FALSE, not a formula or numeric flag.')

def _integer(value,sheet,row,column,maximum=3000):
    if type(value) is not int or not 1<=value<=maximum:_error(sheet,row,column,f'Use a whole number from 1 to {maximum}.')
    return value

def _extras(value,allowed,sheet,row,target):
    extras=_json(value,sheet,row,'extra_json')
    unknown=set(extras)-allowed
    if unknown:_error(sheet,row,'extra_json','Unsupported properties: '+', '.join(sorted(unknown)))
    for key,item in extras.items():
        if key in target:_error(sheet,row,'extra_json',f'{key} is already specified in its column.')
        if key=='excel_row':
            if type(item) is not int or not 1<=item<=100000:_error(sheet,row,'extra_json','excel_row must be a positive integer.')
        elif key=='wizard_context':
            if not isinstance(item,dict) or set(item)!={'version','industry','workflow_type'} or type(item.get('version')) is not int or item['version']!=1 or item.get('workflow_type') not in WORKFLOWS or WORKFLOWS[item['workflow_type']]!=item.get('industry'):_error(sheet,row,'extra_json','Invalid wizard context.')
        elif key in {'room','room_context'}:
            try:_rooms([item])
            except (ValueError,TypeError):_error(sheet,row,'extra_json','Invalid room context.')
        elif key=='topic_sources':
            if not isinstance(item,list) or len(item)>20 or any(not isinstance(x,dict) or set(x)!={'title','url'} or not all(isinstance(v,str) and len(v)<=1000 for v in x.values()) or not x['url'].startswith('https://') for x in item):_error(sheet,row,'extra_json','Invalid descriptive source references.')
        elif not isinstance(item,str) or len(item)>10000:_error(sheet,row,'extra_json',f'{key} must be bounded descriptive text.')
        target[key]=deepcopy(item)

def _package(data):
    if not isinstance(data,bytes) or not data or len(data)>MAX_BYTES:_error('Workbook',0,'file','Provide an .xlsx file up to 5 MB.')
    try:
        archive=zipfile.ZipFile(BytesIO(data));entries=archive.infolist()
        if len(entries)>2000 or sum(x.file_size for x in entries)>MAX_EXPANDED:_error('Workbook',0,'file','Expanded workbook exceeds safety limits.')
        names=[x.filename for x in entries]
        if len(names)!=len(set(names)):_error('Workbook',0,'file','Duplicate package members are not supported.')
        for entry in entries:
            name=entry.filename.lower()
            if entry.flag_bits&1 or '..' in name.split('/') or name.startswith('/') or '\\' in name:_error('Workbook',0,'file','Unsafe workbook package.')
            if any(part in name for part in ['vbaproject','externallinks/','embeddings/','activex/','macrosheets/']):_error('Workbook',0,'file','Macros, embedded objects and external links are not supported.')
            if entry.file_size>MAX_EXPANDED or entry.file_size>max(entry.compress_size,1)*1000:_error('Workbook',0,'file','Excessive compression ratio.')
            if name.endswith(('.xml','.rels')):
                raw=archive.read(entry)
                if b'<!DOCTYPE' in raw.replace(b'\x00',b'').upper() or b'<!ENTITY' in raw.replace(b'\x00',b'').upper():_error('Workbook',0,'file','XML entities are not supported.')
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    # Coordinates, not optional worksheet dimensions, bound expansion.
                    for _,element in ElementTree.iterparse(BytesIO(raw),events=('end',)):
                        tag=element.tag.rsplit('}',1)[-1]
                        if tag=='row':
                            index=element.get('r','')
                            if index and (not index.isascii() or not index.isdigit() or int(index)>100001):_error('Workbook',0,'dimensions','Worksheet row limit exceeded.')
                        elif tag=='c':
                            reference=element.get('r','')
                            match=re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,8})',reference)
                            if reference and not match:_error('Workbook',0,'dimensions','Invalid cell coordinate.')
                            if match:
                                column=0
                                for char in match[1]:column=column*26+ord(char)-64
                                if column>max(map(len,HEADERS.values())) or int(match[2])>100001:_error('Workbook',0,'dimensions','Worksheet coordinate limit exceeded.')
                        element.clear()
                if name.endswith('.rels'):
                    for element in ElementTree.fromstring(raw).iter():
                        if element.get('TargetMode','').lower()=='external':_error('Workbook',0,'file','External relationships are not supported.')
                if name=='[content_types].xml' and b'macroenabled' in raw.lower():_error('Workbook',0,'file','Macro-enabled workbooks are not supported.')
        archive.close()
    except WorkbookValidationError:raise
    except (zipfile.BadZipFile,KeyError,ValueError,ElementTree.ParseError,RuntimeError,NotImplementedError):_error('Workbook',0,'file','Invalid .xlsx package.')

def parse_workbook(data):
    """Read trusted-format bytes without file writes or database access."""
    _package(data)
    from openpyxl import load_workbook
    try:workbook=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    except Exception:_error('Workbook',0,'file','Unable to read this .xlsx workbook.')
    try:
        if workbook.sheetnames==['Checklist']:
            from .checklist_simple_workbook import parse_simple
            return parse_simple(workbook,data)
        tables={};totalrows=0
        for sheet in workbook:
            if sheet.title not in HEADERS:_error(sheet.title,1,'sheet','Unknown worksheet.')
            sheet.reset_dimensions()
            sheet.calculate_dimension(force=True)
            if sheet.max_row>100001 or sheet.max_column>len(HEADERS[sheet.title]):_error(sheet.title,1,'dimensions','Worksheet dimensions exceed format limits.')
            iterator=sheet.iter_rows()
            first=next(iterator,())
            headers=[cell.value for cell in first]
            while headers and headers[-1] is None:headers.pop()
            if not headers:continue
            if len(headers)!=len(set(str(x) for x in headers)) or any(not isinstance(x,str) or x not in HEADERS[sheet.title] for x in headers):_error(sheet.title,1,'headers','Unknown or duplicate header.')
            mandatory=HEADERS[sheet.title][:1] if sheet.title in {'Guide','_Lists'} else HEADERS[sheet.title][:2]
            if any(key not in headers for key in mandatory):_error(sheet.title,1,'headers','Missing reference columns.')
            rows=[]
            for rowno,cells in enumerate(iterator,2):
                totalrows+=1
                if totalrows>100000:_error(sheet.title,rowno,'row','Workbook exceeds 100,000 data rows.')
                values=[]
                for index,cell in enumerate(cells):
                    if cell.data_type in {'f','e'}:_error(sheet.title,rowno,headers[index] if index<len(headers) else str(index+1),'Formulas and Excel errors are not supported.')
                    value=cell.value
                    if value is not None and (not isinstance(value,(str,bool,int,float)) or isinstance(value,float) and not math.isfinite(value)):_error(sheet.title,rowno,str(index+1),'Use plain text, booleans or finite numbers.')
                    if isinstance(value,str) and len(value)>32767:_error(sheet.title,rowno,str(index+1),'Cell text exceeds its limit.')
                    values.append(value)
                if any(value is not None for value in values[len(headers):]):_error(sheet.title,rowno,'headers','Data has no declared column.')
                if any(value not in (None,'') for value in values):rows.append((rowno,dict(zip(headers,values))))
            tables[sheet.title]=rows
        for name in ['Checklists','Sections','Items']:
            if name not in tables:_error(name,1,'sheet','Required worksheet is missing.')
        checklists={};sections={};fields={};orders={};rules_seen=set()
        for row,r in tables['Checklists']:
            sh='Checklists';key=_id(r.get('checklist_key'),sh,row,'checklist_key')
            if key in checklists:_error(sh,row,'checklist_key','Duplicate checklist key.')
            if len(checklists)>=10:_error(sh,row,'checklist_key','At most 10 checklists may be imported.')
            name=_text(r.get('name'),sh,row,'name',160,True);industry=r.get('industry');workflow=r.get('workflow_type')
            if workflow not in WORKFLOWS or WORKFLOWS[workflow]!=industry:_error(sh,row,'workflow_type','Select a supported industry/workflow combination.')
            definition={'sections':[]}
            if r.get('description') not in (None,''):definition['description']=_text(r['description'],sh,row,'description')
            _extras(r.get('extra_json'),ROOT_EXTRA,sh,row,definition)
            if definition.get('wizard_context') and (definition['wizard_context']['industry'],definition['wizard_context']['workflow_type'])!=(industry,workflow):_error(sh,row,'extra_json','Wizard context must match checklist industry and workflow.')
            standard=r.get('standard_key');stages=_json(r.get('stages'),sh,row,'stages',list)
            if standard not in (None,''):
                match=next((x for x in CATALOG if x[0]==standard),None)
                if not match or (match[2],match[3])!=(industry,workflow):_error(sh,row,'standard_key','Standard must match industry and workflow.')
                if standard=='construction_inspection_all':
                    if not stages or any(not isinstance(x,str) or x not in STAGES for x in stages) or len(set(stages))!=len(stages):_error(sh,row,'stages','Use distinct supported construction stages as a JSON array.')
                elif stages:_error(sh,row,'stages','Stages apply only to construction standards.')
                definition['standard_catalog']={'key':standard,'version':1,'industry':industry,'workflow_type':workflow,'stages':stages,'rooms':[]}
            elif stages:_error(sh,row,'stages','Stages require a known construction standard.')
            checklists[key]={'key':key,'name':name,'industry':industry,'workflow_type':workflow,'definition':definition}
        if not checklists:_error('Checklists',2,'checklist_key','Provide at least one checklist.')
        def parent(r,sh,row):
            key=r.get('checklist_key')
            if key not in checklists:_error(sh,row,'checklist_key','Unresolved checklist reference.')
            return key,checklists[key]['definition']
        for row,r in tables['Sections']:
            sh='Sections';key,definition=parent(r,sh,row);sid=_id(r.get('section_id'),sh,row,'section_id')
            if (key,sid) in sections:_error(sh,row,'section_id','Duplicate section ID.')
            section={'id':sid,'title':_text(r.get('title'),sh,row,'title',160,True),'fields':[]}
            for prop in ['group','description','purpose']:
                if r.get(prop) not in (None,''):section[prop]=_text(r[prop],sh,row,prop,160 if prop=='group' else 10000)
            _extras(r.get('extra_json'),SECTION_EXTRA,sh,row,section)
            sections[key,sid]=section;definition['sections'].append(section)
        for row,r in tables['Items']:
            sh='Items';key,definition=parent(r,sh,row);sid=r.get('section_id');fid=_id(r.get('field_id'),sh,row,'field_id')
            if (key,sid) not in sections:_error(sh,row,'section_id','Unresolved section reference.')
            if (key,fid) in fields:_error(sh,row,'field_id','Duplicate field ID within checklist.')
            if len(fields)>=3000:_error(sh,row,'field_id','At most 3000 total fields may be imported.')
            kind=r.get('type')
            if kind not in FIELD_TYPES:_error(sh,row,'type','Unsupported field type.')
            field={'id':fid,'label':_text(r.get('label'),sh,row,'label',3000,True),'type':kind}
            for prop in ['required','photo_required']:
                if r.get(prop) not in (None,''):field[prop]=_bool(r[prop],sh,row,prop)
            for prop in ['system','description']:
                if r.get(prop) not in (None,''):field[prop]=_text(r[prop],sh,row,prop)
            if field.get('system') and (field['system']!='site_id' or kind!='text'):_error(sh,row,'system','Only a text Site ID may be system-linked.')
            if any(r.get(prop) not in (None,'') for prop in ['room_type','room_name','floor']):
                room={'type':r.get('room_type'),'name':r.get('room_name'),'floor':r.get('floor') or ''}
                try:field['room_context']=_rooms([room])[0]
                except (ValueError,TypeError):_error(sh,row,'room_name','Invalid room context.')
            _extras(r.get('extra_json'),FIELD_EXTRA,sh,row,field)
            fields[key,fid]=field;sections[key,sid]['fields'].append(field)
        for sh in ['Choices','TableColumns']:
            for row,r in tables.get(sh,[]):
                key,_=parent(r,sh,row);fid=r.get('field_id');field=fields.get((key,fid))
                if not field:_error(sh,row,'field_id','Unresolved field reference.')
                wanted='select' if sh=='Choices' else 'table'
                if field['type']!=wanted:_error(sh,row,'field_id',f'This row requires a {wanted} field.')
                order=_integer(r.get('order'),sh,row,'order',100 if wanted=='select' else 20)
                group=(sh,key,fid);bucket=orders.setdefault(group,{})
                if order in bucket:_error(sh,row,'order','Duplicate order within field.')
                if wanted=='select':value=_text(r.get('choice'),sh,row,'choice',200,True)
                else:
                    value={'id':_id(r.get('column_id'),sh,row,'column_id'),'label':_text(r.get('label'),sh,row,'label',3000,True)}
                    _extras(r.get('extra_json'),COLUMN_EXTRA,sh,row,value)
                bucket[order]=value
        for (sh,key,fid),bucket in orders.items():
            if sorted(bucket)!=list(range(1,len(bucket)+1)):_error(sh,1,'order','Orders must start at 1 with no gaps.')
            values=[bucket[i] for i in sorted(bucket)]
            identities=values if sh=='Choices' else [x['id'] for x in values]
            if len(set(identities))!=len(identities):_error(sh,1,'choice' if sh=='Choices' else 'column_id','Duplicate choice or column ID.')
            fields[key,fid]['options' if sh=='Choices' else 'columns']=values
        for row,r in tables.get('EvidenceRules',[]):
            sh='EvidenceRules';key,_=parent(r,sh,row);fid=r.get('field_id')
            if (key,fid) not in fields:_error(sh,row,'field_id','Unresolved field reference.')
            if (key,fid) in rules_seen:_error(sh,row,'field_id','Duplicate evidence-rule row.')
            rules_seen.add((key,fid));rules={}
            for prop in RULE_COLUMNS:
                value=r.get(prop)
                if value in (None,''):continue
                if prop in {'required_subjects','any_subjects','required_text','forbidden_subjects'}:value=_json(value,sh,row,prop,list)
                elif prop.startswith('require_'):value=_bool(value,sh,row,prop)
                elif prop in {'version','min_photos'}:value=_integer(value,sh,row,prop,1 if prop=='version' else 20)
                rules[prop]=value
            try:validate_evidence_rules(rules)
            except ValueError as error:_error(sh,row,'rules',str(error))
            fields[key,fid]['evidence_rules']=rules
        for row,r in tables.get('Rooms',[]):
            sh='Rooms';key,definition=parent(r,sh,row);catalog=definition.get('standard_catalog')
            if not catalog or catalog['key']!='construction_inspection_all':_error(sh,row,'checklist_key','Rooms require the construction standard.')
            catalog['rooms'].append({prop:r.get(prop) or '' for prop in ['type','name','floor']})
            try:_rooms(catalog['rooms'])
            except (ValueError,TypeError) as error:_error(sh,row,'name',str(error))
        for item in checklists.values():
            try:validate_definition(item['definition'])
            except ValidationError as error:_error('Checklists',next(row for row,r in tables['Checklists'] if r['checklist_key']==item['key']),'definition','; '.join(error.messages))
        return {'checklists':list(checklists.values()),'warnings':[],'sha256':hashlib.sha256(data).hexdigest()}
    finally:workbook.close()


def workbook_rows(checklists):
    """Return deterministic cell data; caller authors the actual downloadable XLSX.

    Each input is {key,name,industry,workflow_type,definition}. Unsupported source
    properties fail rather than being silently discarded. No operational data.
    """
    result={name:{'headers':list(headers),'rows':[]} for name,headers in HEADERS.items()}
    def add(sheet,values):result[sheet]['rows'].append([values.get(key) for key in HEADERS[sheet]])
    def extras(source,known,allowed,sheet):
        unknown=set(source)-set(known)-allowed
        if unknown:_error(sheet,0,'extra_json','Unsupported source properties: '+', '.join(sorted(unknown)))
        data={key:deepcopy(value) for key,value in source.items() if key in allowed and (key not in known or value=='')}
        return json.dumps(data,ensure_ascii=False,separators=(',',':')) if data else None
    if not isinstance(checklists,list) or not 1<=len(checklists)<=10:_error('Checklists',0,'checklist_key','Provide 1 to 10 checklist records.')
    for item in checklists:
        definition=deepcopy(item['definition']);key=item['key'];catalog=definition.get('standard_catalog',{})
        try:validate_definition(definition)
        except ValidationError as error:_error('Checklists',0,'definition','; '.join(error.messages))
        wizard=definition.get('wizard_context')
        if wizard and (wizard.get('industry'),wizard.get('workflow_type'))!=(item['industry'],item['workflow_type']):_error('Checklists',0,'extra_json','Wizard context must match checklist industry and workflow.')
        if catalog:
            known=next((entry for entry in CATALOG if entry[0]==catalog.get('key')),None)
            if set(catalog)!={'key','version','industry','workflow_type','stages','rooms'} or catalog.get('version')!=1 or not known or (catalog.get('industry'),catalog.get('workflow_type'))!=(item['industry'],item['workflow_type']) or (known[2],known[3])!=(item['industry'],item['workflow_type']):_error('Checklists',0,'standard_key','Unsupported standard metadata; it cannot be discarded on export.')
        add('Checklists',{'checklist_key':key,'name':item['name'],'industry':item['industry'],'workflow_type':item['workflow_type'],'description':definition.get('description') or None,'standard_key':catalog.get('key'),'stages':json.dumps(catalog.get('stages',[])) if catalog else None,'extra_json':extras(definition,{'sections','standard_catalog','description'},ROOT_EXTRA,'Checklists')})
        for section in definition['sections']:
            add('Sections',{'checklist_key':key,'section_id':section['id'],'title':section['title'],'group':section.get('group') or None,'description':section.get('description') or None,'purpose':section.get('purpose') or None,'extra_json':extras(section,{'id','title','fields','group','description','purpose'},SECTION_EXTRA,'Sections')})
            for field in section['fields']:
                room=field.get('room_context',{})
                add('Items',{'checklist_key':key,'section_id':section['id'],'field_id':field['id'],'label':field['label'],'type':field['type'],**{prop:field.get(prop) for prop in ['required','photo_required','system']},'description':field.get('description') or None,'room_type':room.get('type'),'room_name':room.get('name'),'floor':room.get('floor'),'extra_json':extras(field,{'id','label','type','required','photo_required','system','description','options','columns','evidence_rules','room_context'},FIELD_EXTRA,'Items')})
                for order,choice in enumerate(field.get('options',[]),1):add('Choices',{'checklist_key':key,'field_id':field['id'],'choice':choice,'order':order})
                for order,column in enumerate(field.get('columns',[]),1):add('TableColumns',{'checklist_key':key,'field_id':field['id'],'column_id':column['id'],'label':column['label'],'order':order,'extra_json':extras(column,{'id','label'},COLUMN_EXTRA,'TableColumns')})
                if 'evidence_rules' in field:add('EvidenceRules',{'checklist_key':key,'field_id':field['id'],**{prop:json.dumps(value,ensure_ascii=False) if isinstance(value,list) else value for prop,value in field['evidence_rules'].items()}})
        for room in catalog.get('rooms',[]):add('Rooms',{'checklist_key':key,**room})
    guide=[
        ('Purpose and approval','Workbook format version 1. Import creates new unpublished drafts only after preview and explicit commit; preview expires after 30 minutes. It does not publish, assign tasks, run AI, enter survey answers or grant approvals. Existing published revisions and survey snapshots remain unchanged.'),
        ('Limits','Maximum 10 checklists and 3000 total item fields; upload at most 5 MB .xlsx, expanded package at most 50 MB. No macros, formulas, external links, embedded objects, unsupported sheets, unsupported conditions or unknown columns.'),
        ('How rows connect','Use a stable checklist_key starting with a letter, followed by letters, digits, underscores or hyphens, maximum 80 characters. section_id is unique within a checklist; field_id is unique across its sections. All child references must resolve. Section/item row order is preserved. Choice and table-column orders start at 1 without gaps or duplicates.'),
        ('Industry and workflow','Supported pairs only: telecom / telecom_survey; telecom / telecom_inspection; construction / construction_inspection. Template industry and workflow must match project and task. Industry does not itself establish technical qualification.'),
        ('Text and textarea','text: single-line free text. textarea: longer narrative. Neither supports executable calculations or conditional visibility. Descriptive unit/equipment_type/operator/excel_row metadata may be preserved through allowed extra_json keys.'),
        ('Number and date','number: numeric text checked by the application; no workbook-authored numeric minimum/maximum, units conversion or automatic acceptance. date: date entry in the application. Workbook cells define the fields, not their future answers.'),
        ('Select','select requires 1 to 100 unique Choices rows for its field_id. choice is exact display text, maximum 200 characters; order defines display order. No cascading or conditional choices are supported.'),
        ('Inspection','inspection supports Checked, Defect found or Not applicable decisions entered by the assigned human in the application. Not applicable requires a reason. This file cannot prefill decisions or approve work.'),
        ('Table','table requires 1 to 20 unique TableColumns IDs with labels and consecutive order. These are repeated-row column definitions only. Test results, measurements, report IDs, dates, human pass/fail and retest rows are entered later in the application, not imported as checklist answers.'),
        ('Photo','photo defines a photo evidence location. Files are uploaded in the application. required and photo_required are independent TRUE/FALSE flags: required requires a value/evidence appropriate to type; photo_required requires linked evidence. Neither means AI or engineering approval.'),
        ('Optional cells and extra_json','Blank optional columns mean absent, using existing application defaults. FALSE is explicit false. extra_json preserves only supported descriptive properties, not conditions, roles, answers, approvals or database IDs. Do not duplicate a property in both its normal column and extra_json. Exported empty descriptive properties may appear in extra_json to preserve exact values.'),
        ('Evidence rules format','At most one EvidenceRules row per field. version is integer 1 when specified. JSON array columns use e.g. ["tower","ladder"], never pipe-separated text. All listed strings are literal evidence descriptions/text, never executable patterns. Blank rule cells mean no declared requirement; no rule changes an engineering result.'),
        ('Subject rules','required_subjects: all listed visible subjects; any_subjects: at least one listed subject; forbidden_subjects: listed contradictory subjects. Each list has at most 12 distinct, nonempty strings of at most 100 characters. Ambiguous or missing support remains subject to review.'),
        ('Text and identity rules','required_text: JSON list of literal text phrases, same 12-entry/100-character limits. expected_site_id: literal Site ID, 2 to 64 letters/digits/dot/underscore/hyphen starting with a letter or digit; no regex. require_readable_text: TRUE/FALSE. Text is evidence data, never instructions; dates, GPS and identity are not guessed from captions.'),
        ('Capture and quality rules','require_qr, require_gps and require_capture_time: TRUE/FALSE. minimum_quality: any or clear. These requirements rely on saved source-bound evidence assessments; missing support blocks declared sufficiency. They do not create measurements or compliance decisions.'),
        ('Photo count and notes','min_photos: integer 1 to 20, distinct human-confirmed source photos required; this is the supported minimum count, not a numerical engineering acceptance limit. notes: descriptive text at most 1000 characters. No maximum-photo acceptance condition is supported.'),
        ('Standard metadata and stages','standard_key is optional and must be a known catalog key matching industry/workflow. stages is a JSON array of distinct supported construction stage codes; see _Lists. No stages on telecom standards. Legacy stage codes remain supported. A new draft can be edited and qualified for its project; stored published sources are never rewritten.'),
        ('Rooms','Construction standard Rooms rows use a supported type plus room name and optional floor; maximum 40 instances. Name maximum 80, floor maximum 60 characters. Name/floor pairs must be unique ignoring case. Repeated room types are separate sections and field IDs. Item room_type/room_name/floor preserve exact room context; they are not automatic room identity proof.'),
        ('Structural activity registers','Granular structural activities end with a required repeated test-report register. Enter actual report-backed rows later, or an explicit N/A row with reason and approved inspection-plan applicability reference. General legacy stage registers are optional. There are no invented numerical pass/fail thresholds.'),
        ('Preview and correction','Validate every sheet and row in preview before commit. Errors identify sheet, row and column. Correct and re-upload; no partial import is committed. A successful commit creates drafts, after which the existing edit, technical qualification, publication and assignment controls apply.'),
    ]
    for topic,instructions in guide:add('Guide',{'topic':topic,'instructions':instructions})
    lists={'field_type':sorted(FIELD_TYPES),'industry':['telecom','construction'],'workflow_type':list(WORKFLOWS),'boolean':['TRUE','FALSE'],'minimum_quality':['any','clear'],'room_type':list(ROOM_TYPES),'construction_stage':list(STAGES)}
    for kind,values in lists.items():
        for value in values:add('_Lists',{'list':kind,'value':value})
    for stage in STAGES:add('Guide',{'topic':'Construction stage: '+stage,'instructions':STAGE_LABELS.get(stage,stage.replace('_',' ').title())})
    return result


def export_workbook(checklists):
    """Runtime XLSX download using the installed stack, entirely in memory.

    All strings are explicitly typed as data, including leading '=' text.
    Dropdowns assist entry only; parse_workbook remains authoritative.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.utils import get_column_letter
    tables=workbook_rows(checklists)
    enums={'field_type':sorted(FIELD_TYPES),'boolean':['TRUE','FALSE'],'industry':['telecom','construction'],'workflow':list(WORKFLOWS),'room_type':list(ROOM_TYPES),'quality':['any','clear'],'construction_stage':list(STAGES)}
    tables['_Lists']['rows']=[[kind,value] for kind,values in enums.items() for value in values]
    workbook=Workbook();workbook.remove(workbook.active)
    for name,table in tables.items():
        sheet=workbook.create_sheet(name)
        for row_index,values in enumerate([table['headers'],*table['rows']],1):
            for column,value in enumerate(values,1):
                cell=sheet.cell(row_index,column,value)
                if isinstance(value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if row_index==1:
                    cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True)
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for column,header in enumerate(table['headers'],1):sheet.column_dimensions[get_column_letter(column)].width=50 if header in {'description','instructions','extra_json','label','notes'} else 24
    ranges={};start=2
    for kind,values in enums.items():
        address=f"'_Lists'!$B${start}:$B${start+len(values)-1}"
        ranges[kind]='choices_'+kind
        workbook.defined_names.add(DefinedName(ranges[kind],attr_text=address))
        start+=len(values)
    targets={'Items':{'type':'field_type','required':'boolean','photo_required':'boolean','room_type':'room_type'},'Checklists':{'industry':'industry','workflow_type':'workflow'},'Rooms':{'type':'room_type'},'EvidenceRules':{'require_readable_text':'boolean','require_qr':'boolean','require_gps':'boolean','require_capture_time':'boolean','minimum_quality':'quality'}}
    for name,columns in targets.items():
        sheet=workbook[name]
        for header,kind in columns.items():
            validation=DataValidation(type='list',formula1=ranges[kind],allow_blank=True)
            validation.errorTitle='Choose a supported value';validation.error='Select a listed value or leave optional cells blank.';validation.showErrorMessage=True;validation.errorStyle='stop'
            sheet.add_data_validation(validation);letter=get_column_letter(HEADERS[name].index(header)+1)
            validation.add(f'{letter}2:{letter}{max(sheet.max_row+100,1002)}')
    workbook['_Lists'].sheet_state='hidden'
    output=BytesIO();workbook.save(output);workbook.close()
    data=output.getvalue()
    # Apply the same contract to runtime downloads before returning bytes.
    parse_workbook(data)
    return data
