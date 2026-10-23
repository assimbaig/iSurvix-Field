"""Independent HTTP acceptance for category capture followed by immediate QA.

All records, media and model access are isolated by scripts/run_backend_tests.py.
Synthetic job states exercise asynchronous boundaries, not native AI accuracy.
"""
from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from PIL import Image
from openpyxl import load_workbook

from . import test_survey_processing as fixtures, category_intake, survey_processing as pipeline
from .independent_workflow import configure_policy
from .models import (Checklist, ChecklistTemplate, PhotoDraft, PhotoEvidence, PhotoJob, PhotoPlacement,
    Profile, WorkAssignment, WorkflowTemplateQualification)
from .review_route_service import route_context
from .roadmap_service import definition_digest
from .survey_processing_models import SurveyProcessingJob, SurveyProcessingResult, SurveyProcessingRevision
from .field_photo_service import schema_digest
from .template_service import missing_required


@override_settings(SECRET_KEY='isolated-category-submission-acceptance',
    ROOT_URLCONF='checklist.urls_local', ALLOWED_HOSTS=['testserver'], SECURE_SSL_REDIRECT=False,
    PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class EngineerCategorySubmissionAcceptanceTests(TestCase):
    person = fixtures.SurveyProcessingTests.person
    add_photo = fixtures.SurveyProcessingTests.add_photo
    analysed = fixtures.SurveyProcessingTests.analysed
    processing = fixtures.SurveyProcessingTests.processing

    def setUp(self):
        fixtures.SurveyProcessingTests.setUp(self)
        self.schema = deepcopy(self.schema)
        self.schema['sections'][0]['fields'][2]['required'] = True
        self.template = ChecklistTemplate.objects.create(name='SYNTHETIC category capture acceptance',
            definition=self.schema, published=True, checklist_required=True)
        self.project.checklist_template = self.template
        self.project.review_layers = 4
        self.project.save(update_fields=['checklist_template', 'review_layers'])
        self.checklist = Checklist.objects.create(project=self.project, user=self.engineer,
            quality_reviewer=self.quality, site_id='SYNTHETIC-CATEGORY-ACCEPTANCE', answer_data={},
            coordinate_reference={'source': 'Explicit synthetic location', 'crs': 'EPSG:4326'})
        WorkflowTemplateQualification.objects.create(project=self.project, template=self.template,
            workflow_type=self.checklist.workflow_type, definition_digest=definition_digest(self.schema),
            criteria_reference='SYNTHETIC label criteria', inspection_method='Inspect source evidence',
            reviewed_by=self.lead, reason='Isolated category capture acceptance')
        configure_policy(self.admin, self.checklist, self.quality, self.lead,
            'Synthetic category review route', customer_reviewer=self.customer,
            customer_required=True, enable_route=True)
        self.assignment = WorkAssignment.objects.create(project=self.project, assigned_by=self.lead,
            assigned_to=self.engineer, checklist=self.checklist, site_id=self.checklist.site_id,
            latitude=24, longitude=46, description='SYNTHETIC independent category capture')
        self.client.force_login(self.engineer)

    def post(self, name, payload, actor=None, **kwargs):
        if actor is not None:
            self.client.force_login(actor)
        return self.client.post(reverse(name, args=[self.checklist.pk]),
            json.dumps(payload), content_type='application/json', **kwargs)

    def state(self, actor=None):
        if actor is not None:
            self.client.force_login(actor)
        response = self.client.get(reverse('photo_capture_state', args=[self.checklist.pk]))
        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertIsInstance(payload['digest'], str)
        self.assertEqual(len(payload['digest']), 64)
        return payload

    def raw(self, label):
        color = tuple(hashlib.sha256(label.encode()).digest()[:3])
        image = BytesIO()
        Image.new('RGB', (65, 43), color).save(image, 'PNG')
        return image.getvalue()

    def upload(self, group='Civil', labels=('synthetic-one',), actor=None, finish=True):
        if actor is not None:
            self.client.force_login(actor)
        session = self.post('photo_processing_upload_session', {'action': 'start'})
        self.assertEqual(session.status_code, 200, session.content)
        token = session.json()['token']
        files = [SimpleUploadedFile(label + '.png', self.raw(label), content_type='image/png')
                 for label in labels]
        response = self.client.post(reverse('photo_upload', args=[self.checklist.pk]),
            {'group': group, 'processing_upload_token': token, 'files': files})
        self.assertEqual(response.status_code, 201, response.content)
        if finish:
            closed = self.post('photo_processing_upload_session', {'action': 'finish', 'token': token})
            self.assertEqual(closed.status_code, 200, closed.content)
        return response.json()['photos'], token

    def submit(self, digest=None, actor=None):
        if actor is not None:
            self.client.force_login(actor)
        digest = digest or self.state()['digest']
        response = self.post('photo_capture_submit', {'uploads_finished': True, 'expected_digest': digest})
        return response, digest

    def assert_submitted(self):
        self.checklist.refresh_from_db()
        self.assignment.refresh_from_db()
        self.assertEqual(self.checklist.review_route_record.stage, 'quality_engineer')
        self.assertEqual(self.checklist.status, Checklist.Status.SUBMITTED)
        self.assertEqual(self.assignment.status, WorkAssignment.Status.SUBMITTED)
        self.assertTrue(missing_required(self.checklist), 'Engineer capture must not manufacture missing answers')
        return self.checklist.processing_revisions.get()

    def qa_state(self, actor=None):
        self.client.force_login(actor or self.quality)
        response = self.client.get(reverse('photo_capture_review_state', args=[self.checklist.pk]))
        self.assertEqual(response.status_code, 200, response.content)
        return response.json()

    def qa_action(self, action, **data):
        state = self.qa_state()
        return self.post('photo_capture_review_action', {'action': action,
            'expected_review_digest': state['review_digest'], **data})

    def route(self, actor, action, reason='Synthetic source-backed review'):
        self.checklist.refresh_from_db()
        token = route_context(actor, self.checklist)['token']
        self.client.force_login(actor)
        return self.client.post(reverse('review_route_action', args=[self.checklist.pk]),
            {'action': action, 'reason': reason, 'expected_token': token})

    def complete_qa_manually(self):
        for ident, value in [('site_id', self.checklist.site_id),
                             ('contact', 'Synthetic QA contact'), ('manufacturer', 'SyntheticCo')]:
            result = self.qa_action('answer', field_id=ident, value=value,
                reason='Synthetic recorded source checked independently by QA',
                photo_ids=[self.checklist.photo_evidence.get(filename='synthetic-one.png').pk]
                    if ident == 'manufacturer' else [])
            self.assertEqual(result.status_code, 200, result.content)
        completed = self.qa_action('complete')
        self.assertEqual(completed.status_code, 200, completed.content)
        self.assertTrue(completed.json()['completed'])
        return completed.json()

    def test_multiple_category_uploads_preserve_originals_and_only_submit_starts_async_ai(self):
        expected = {}
        for category in ('Civil', 'Electrical', 'Tower', 'General Info'):
            labels = (category + '-one', category + '-two')
            rows, _ = self.upload(category, labels)
            for row, label in zip(rows, labels):
                expected[row['id']] = ('' if category == 'General Info' else category, self.raw(label))
        self.assertEqual(self.checklist.photo_evidence.count(), 8)
        self.assertFalse(self.checklist.processing_revisions.exists())
        self.assertFalse(PhotoJob.objects.filter(checklist=self.checklist).exists())
        for photo in self.checklist.photo_evidence.all():
            group, raw = expected[photo.pk]
            self.assertEqual(photo.group, group)
            self.assertEqual(Path(photo.original.path).read_bytes(), raw)
            self.assertEqual(photo.sha256, hashlib.sha256(raw).hexdigest())
            self.assertNotEqual(photo.preview.name, photo.original.name)
        response, _ = self.submit()
        self.assertIn(response.status_code, (200, 201), response.content)
        revision = self.assert_submitted()
        self.assertEqual(revision.processing_job.status, 'queued')
        self.assertEqual(revision.snapshot['category_intake_version'], 1)
        self.assertEqual({p['id']: p['group'] for p in revision.snapshot['photos']},
                         {pk: value[0] for pk, value in expected.items()})
        self.assertEqual(self.checklist.submission_revisions.count(), 1)
        self.assertFalse(PhotoPlacement.objects.filter(photo__checklist=self.checklist).exists())
        self.assertEqual(self.checklist.answer_data, {})

    def test_duplicate_upload_and_repeated_submit_keep_one_job_and_frozen_revision(self):
        first, _ = self.upload('Civil')
        repeated, _ = self.upload('Civil')
        self.assertEqual(first[0]['id'], repeated[0]['id'])
        self.assertTrue(repeated[0]['duplicate'])
        response, signature = self.submit()
        self.assertIn(response.status_code, (200, 201), response.content)
        revision = self.assert_submitted()
        original_snapshot = deepcopy(revision.snapshot)
        second, _ = self.submit(signature)
        self.assertIn(second.status_code, (200, 201), second.content)
        revision.refresh_from_db()
        self.assertEqual(revision.snapshot, original_snapshot)
        self.assertEqual(self.checklist.processing_revisions.count(), 1)
        self.assertEqual(self.checklist.submission_revisions.count(), 1)
        self.assertEqual(SurveyProcessingJob.objects.filter(revision__checklist=self.checklist).count(), 1)

    def test_category_change_remove_and_preview_work_before_submission_then_freeze(self):
        rows, _ = self.upload('Civil', ('move-this', 'remove-this'))
        moving, removing = [PhotoEvidence.objects.get(pk=row['id']) for row in rows]
        moved = self.post('photo_capture_category', {'photo_id': moving.pk, 'group': 'Tower',
            'expected_digest': self.state()['digest']})
        self.assertEqual(moved.status_code, 200, moved.content)
        moving.refresh_from_db()
        self.assertEqual(moving.group, 'Tower')
        self.assertEqual(Path(moving.original.path).read_bytes(), self.raw('move-this'))
        for variant in ('preview', 'original'):
            preview = self.client.get(reverse('photo_evidence', args=[self.checklist.pk, moving.pk, variant]))
            self.assertEqual(preview.status_code, 200)
            content = b''.join(preview.streaming_content)
            self.assertTrue(content)
            if variant == 'original':
                self.assertEqual(content, self.raw('move-this'))
            preview.close()
        removed = self.post('photo_capture_remove', {'photo_id': removing.pk,
            'expected_digest': self.state()['digest']})
        self.assertEqual(removed.status_code, 200, removed.content)
        self.assertNotIn(removing.pk, [p['id'] for p in self.state()['photos']])
        response, _ = self.submit()
        self.assertIn(response.status_code, (200, 201), response.content)
        revision = self.assert_submitted()
        snapshot = deepcopy(revision.snapshot)
        for name, extra in (('photo_capture_category', {'group': 'Electrical'}), ('photo_capture_remove', {})):
            denied = self.post(name, {'photo_id': moving.pk, 'expected_digest': self.state()['digest'], **extra})
            self.assertIn(denied.status_code, (403, 409), denied.content)
        revision.refresh_from_db()
        self.assertEqual(revision.snapshot, snapshot)
        self.assertEqual(Path(moving.original.path).read_bytes(), self.raw('move-this'))

    def test_active_upload_and_corrupt_batch_cannot_be_silently_submitted_or_discarded(self):
        _, token = self.upload('Civil', ('good-upload',), finish=False)
        blocked, _ = self.submit()
        self.assertEqual(blocked.status_code, 409, blocked.content)
        self.assertFalse(self.checklist.submission_revisions.exists())
        self.assertFalse(self.checklist.processing_revisions.exists())
        self.assertEqual(self.checklist.photo_evidence.count(), 1)
        self.assertEqual(self.post('photo_processing_upload_session', {'action': 'finish', 'token': token}).status_code, 200)
        _, token = self.upload('Tower', ('prior-good',), finish=False)
        corrupt = self.client.post(reverse('photo_upload', args=[self.checklist.pk]), {
            'group': 'Tower', 'processing_upload_token': token,
            'files': [SimpleUploadedFile('valid.png', self.raw('atomic-valid'), content_type='image/png'),
                      SimpleUploadedFile('bad.jpg', b'not an image', content_type='image/jpeg')]})
        self.assertEqual(corrupt.status_code, 400, corrupt.content)
        self.assertEqual(self.checklist.photo_evidence.count(), 2)
        self.assertEqual(self.post('photo_processing_upload_session', {'action': 'finish', 'token': token}).status_code, 200)
        self.upload('Tower', ('atomic-valid',))
        completed, _ = self.submit()
        self.assertIn(completed.status_code, (200, 201), completed.content)
        self.assert_submitted()

    def test_stale_category_edits_and_foreign_photo_ids_are_rejected(self):
        rows, _ = self.upload('Civil')
        old = self.state()['digest']
        first = self.post('photo_capture_category', {'photo_id': rows[0]['id'], 'group': 'Electrical',
            'expected_digest': old})
        self.assertEqual(first.status_code, 200)
        stale = self.post('photo_capture_category', {'photo_id': rows[0]['id'], 'group': 'Tower',
            'expected_digest': old})
        self.assertEqual(stale.status_code, 409, stale.content)
        self.assertEqual(self.checklist.photo_evidence.get().group, 'Electrical')
        foreign = self.photo  # Base fixture belongs to a different saved survey.
        for name in ('photo_capture_category', 'photo_capture_remove'):
            denied = self.post(name, {'photo_id': foreign.pk, 'group': 'Civil',
                'expected_digest': self.state()['digest']})
            self.assertIn(denied.status_code, (403, 404), denied.content)
        invalid = self.post('photo_capture_category', {'photo_id': rows[0]['id'], 'group': 'Fabricated category',
            'expected_digest': self.state()['digest']})
        self.assertEqual(invalid.status_code, 409, invalid.content)

    def test_wrong_roles_peer_and_csrf_cannot_submit_or_reorganize_engineer_capture(self):
        rows, _ = self.upload('Civil')
        signature = self.state()['digest']
        peer = self.person('category_peer', Profile.Roles.ENGINEER)
        super_admin = User.objects.create_superuser('category_platform_admin', password=None)
        for actor in (self.admin, self.quality, self.lead, self.customer, peer, super_admin):
            with self.subTest(actor=actor.profile.role):
                self.client.force_login(actor)
                for name, extra in (('photo_capture_submit', {'uploads_finished': True}),
                                    ('photo_capture_category', {'photo_id': rows[0]['id'], 'group': 'Tower'}),
                                    ('photo_capture_remove', {'photo_id': rows[0]['id']})):
                    denied = self.post(name, {'expected_digest': signature, **extra})
                    self.assertIn(denied.status_code, (403, 404), denied.content)
        self.assertFalse(self.checklist.processing_revisions.exists())
        self.assertEqual(self.checklist.photo_evidence.get().group, 'Civil')
        strict = Client(enforce_csrf_checks=True)
        strict.force_login(self.engineer)
        denied = strict.post(reverse('photo_capture_submit', args=[self.checklist.pk]),
            json.dumps({'uploads_finished': True, 'expected_digest': signature}), content_type='application/json')
        self.assertEqual(denied.status_code, 403)

    def test_engineer_cannot_mutate_ai_matches_answers_or_reenter_review_after_handoff(self):
        self.upload('Civil')
        response, _ = self.submit()
        self.assertIn(response.status_code, (200, 201), response.content)
        revision = self.assert_submitted()
        for decision in ('accept', 'reject', 'engineer_complete', 'record', 'geometry'):
            denied = self.client.post(reverse('survey_processing_action', args=[self.checklist.pk]),
                {'action': 'review', 'revision_id': revision.pk, 'decision': decision,
                 'finding_ids': json.dumps(['manufacturer']), 'reason': 'Must be QA-owned'})
            self.assertIn(denied.status_code, (403, 409), denied.content)
        mapping = self.client.post(reverse('photo_mapping', args=[self.checklist.pk]),
            json.dumps({'action': 'manual', 'photo_id': self.checklist.photo_evidence.get().pk,
                        'field_id': 'manufacturer'}), content_type='application/json')
        self.assertIn(mapping.status_code, (403, 409), mapping.content)
        saved = self.client.post(reverse('custom_checklist_save', args=[self.checklist.pk]),
            json.dumps({'answers': {'contact': 'Cannot edit frozen capture'}}), content_type='application/json')
        self.assertIn(saved.status_code, (403, 409), saved.content)
        self.assertFalse(revision.processing_events.filter(action='engineer_complete').exists())

    def test_ai_state_never_returns_queued_running_or_failed_submission_to_engineer(self):
        self.upload('Civil')
        response, _ = self.submit()
        self.assertIn(response.status_code, (200, 201), response.content)
        revision = self.assert_submitted()
        for status in ('queued', 'analysing_evidence', 'failed'):
            with self.subTest(ai_status=status):
                job = revision.processing_job
                job.status = status
                job.error = 'Synthetic model unavailable' if status == 'failed' else ''
                job.save(update_fields=['status', 'error'])
                self.client.force_login(self.quality)
                qa = self.client.get(reverse('custom_checklist', args=[self.checklist.pk]))
                self.assertEqual(qa.status_code, 200, qa.content)
                self.checklist.refresh_from_db()
                self.assertEqual(self.checklist.review_route_record.stage, 'quality_engineer')
                self.client.force_login(self.engineer)
                self.assertEqual(self.client.get(reverse('photo_workspace', args=[self.checklist.pk])).status_code, 200)
                self.assertFalse(revision.processing_events.filter(action='engineer_complete').exists())

    def test_same_project_unassigned_reviewers_cannot_read_private_category_package(self):
        self.upload('Civil')
        submitted, _ = self.submit()
        self.assertIn(submitted.status_code, (200, 201), submitted.content)
        for actor in (self.person('unassigned_quality', Profile.Roles.QUALITY_ENGINEER),
                      self.person('unassigned_lead', Profile.Roles.TEAM_LEAD)):
            self.client.force_login(actor)
            for name in ('photo_capture_state', 'photo_capture_review_state',
                         'survey_processing_state', 'photo_state'):
                with self.subTest(actor=actor.profile.role, route=name):
                    response = self.client.get(reverse(name, args=[self.checklist.pk]))
                    self.assertIn(response.status_code, (403, 404), response.content)

    def test_qa_must_complete_required_review_but_can_do_so_while_ai_is_queued(self):
        self.upload('Civil')
        submitted, _ = self.submit()
        self.assertIn(submitted.status_code, (200, 201), submitted.content)
        revision = self.assert_submitted()
        source_snapshot = deepcopy(revision.snapshot)
        technical = self.checklist.submission_revisions.get()
        engineer_snapshot = deepcopy(technical.snapshot)
        early = self.route(self.quality, 'approve')
        self.assertIn(early.status_code, (302, 409), early.content)
        self.checklist.refresh_from_db()
        self.assertEqual(self.checklist.review_route_record.stage, 'quality_engineer')
        self.assertTrue(self.qa_state()['missing'])
        complete = self.complete_qa_manually()
        self.assertEqual(revision.processing_job.status, 'queued')
        self.assertFalse(complete['missing'])
        revision.refresh_from_db(); technical.refresh_from_db(); self.checklist.refresh_from_db()
        self.assertEqual(revision.snapshot, source_snapshot)
        self.assertEqual(technical.snapshot, engineer_snapshot)
        self.assertEqual(self.checklist.answer_data, {})
        self.assertEqual(revision.processing_events.filter(action='qa_answer', actor=self.quality).count(), 3)
        self.assertFalse(revision.processing_events.filter(action='engineer_complete').exists())
        approved = self.route(self.quality, 'approve')
        self.assertEqual(approved.status_code, 302, approved.content)
        self.checklist.refresh_from_db()
        self.assertEqual(self.checklist.review_route_record.stage, 'team_lead')

    def test_late_ai_cannot_replace_completed_qa_answers_or_frozen_engineer_snapshot(self):
        self.upload('Civil')
        submitted, _ = self.submit()
        self.assertIn(submitted.status_code, (200, 201), submitted.content)
        revision = self.assert_submitted()
        snapshot = deepcopy(revision.snapshot)
        self.complete_qa_manually()
        reviewed_before = deepcopy(category_intake.reviewed_payload(revision))
        late = pipeline.prepare_checklist(snapshot, {'findings': [], 'missing': [], 'conflicts': []})
        for item in late['checklist_items']:
            if item['id'] == 'manufacturer':
                item.update(value='UNVERIFIED LATE MODEL GUESS', proposed_answer='UNVERIFIED LATE MODEL GUESS')
        SurveyProcessingResult.objects.create(revision=revision, payload=late, digest=pipeline.digest(late))
        self.assertEqual(category_intake.reviewed_payload(revision), reviewed_before)
        revision.refresh_from_db()
        self.assertEqual(revision.snapshot, snapshot)
        self.assertEqual(self.route(self.quality, 'approve').status_code, 302)
        self.assertEqual(self.route(self.lead, 'approve').status_code, 302)

    def test_qa_retry_keeps_failed_job_sources_and_engineer_handoff_intact(self):
        self.upload('Civil')
        submitted, _ = self.submit()
        self.assertIn(submitted.status_code, (200, 201), submitted.content)
        revision = self.assert_submitted()
        job = revision.processing_job
        job.status = 'failed'; job.error = 'Synthetic missing local runtime'
        job.save(update_fields=['status', 'error'])
        snapshot = deepcopy(revision.snapshot)
        retry = self.qa_action('retry')
        self.assertEqual(retry.status_code, 200, retry.content)
        second = self.qa_action('retry')
        self.assertEqual(second.status_code, 200, second.content)
        job.refresh_from_db(); revision.refresh_from_db(); self.checklist.refresh_from_db()
        self.assertEqual(job.status, 'queued')
        self.assertEqual(job.error, '')
        self.assertEqual(revision.snapshot, snapshot)
        self.assertEqual(self.checklist.review_route_record.stage, 'quality_engineer')
        self.assertEqual(revision.processing_events.filter(action='qa_retry').count(), 1)
        self.assertEqual(self.checklist.processing_revisions.count(), 1)
        denied = self.post('photo_capture_review_action', {'action': 'retry'}, actor=self.engineer)
        self.assertIn(denied.status_code, (403, 404), denied.content)

    def test_qa_return_allows_new_engineer_revision_while_previous_ai_job_is_unfinished(self):
        self.upload('Civil')
        submitted, _ = self.submit()
        self.assertIn(submitted.status_code, (200, 201), submitted.content)
        first = self.assert_submitted()
        first_snapshot = deepcopy(first.snapshot)
        returned = self.route(self.quality, 'reject', 'Synthetic Engineer must add a missing photograph')
        self.assertEqual(returned.status_code, 302, returned.content)
        self.client.force_login(self.engineer)
        self.upload('Tower', ('synthetic-correction',))
        resubmitted, _ = self.submit()
        self.assertIn(resubmitted.status_code, (200, 201), resubmitted.content)
        first.refresh_from_db(); self.checklist.refresh_from_db()
        latest = self.checklist.processing_revisions.order_by('-number').first()
        self.assertEqual(first.snapshot, first_snapshot)
        self.assertEqual(latest.number, 2)
        self.assertNotEqual(latest.digest, first.digest)
        self.assertEqual(len(latest.snapshot['photos']), 2)
        self.assertEqual(self.checklist.review_route_record.stage, 'quality_engineer')
        self.assertEqual(self.checklist.submission_revisions.count(), 2)
        self.assertEqual(self.checklist.processing_revisions.count(), 2)

    def test_client_cannot_read_removed_pre_submit_source_through_state_or_direct_url(self):
        kept, _ = self.upload('Civil')
        removed, _ = self.upload('Tower', ('removed-source',))
        removed_photo = self.checklist.photo_evidence.get(pk=removed[0]['id'])
        before_bytes = Path(removed_photo.original.path).read_bytes()
        excluded = self.post('photo_capture_remove', {'photo_id': removed_photo.pk,
            'expected_digest': self.state()['digest']})
        self.assertEqual(excluded.status_code, 200, excluded.content)
        submitted, _ = self.submit()
        self.assertIn(submitted.status_code, (200, 201), submitted.content)
        self.complete_qa_manually()
        self.assertEqual(self.route(self.quality, 'approve').status_code, 302)
        self.assertEqual(self.route(self.lead, 'approve').status_code, 302)
        self.client.force_login(self.customer)
        state = self.client.get(reverse('photo_state', args=[self.checklist.pk]))
        self.assertEqual(state.status_code, 200, state.content)
        self.assertEqual([p['id'] for p in state.json()['photos']], [kept[0]['id']])
        denied = self.client.get(reverse('photo_evidence', args=[self.checklist.pk, removed_photo.pk, 'original']))
        self.assertIn(denied.status_code, (403, 404), 'Excluded source must not bypass submitted-snapshot privacy')
        self.assertEqual(Path(removed_photo.original.path).read_bytes(), before_bytes)
        allowed = self.client.get(reverse('photo_evidence', args=[self.checklist.pk, kept[0]['id'], 'original']))
        self.assertEqual(allowed.status_code, 200)
        self.assertEqual(b''.join(allowed.streaming_content), self.raw('synthetic-one'))
        allowed.close()

    def test_existing_excel_export_uses_qa_review_without_rewriting_engineer_answers(self):
        self.upload('Civil')
        submitted, _ = self.submit()
        self.assertIn(submitted.status_code, (200, 201), submitted.content)
        revision = self.assert_submitted()
        immutable_snapshot = deepcopy(revision.snapshot)
        self.complete_qa_manually()
        response = self.client.get(reverse('custom_checklist_download', args=[self.checklist.pk]))
        self.assertEqual(response.status_code, 200, response.content[:200])
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook['Site survey']
        exported = {row[1]: row[2] for row in sheet.iter_rows(values_only=True)
                    if len(row) > 2 and row[1]}
        self.assertEqual(exported['Survey contact'], 'Synthetic QA contact')
        self.assertEqual(exported['Cabinet manufacturer'], 'SyntheticCo')
        for section in self.schema['sections']:
            for field in section['fields']:
                self.assertIn(field['label'], exported)
        self.assertIsNone(exported['Recorded tower height'])
        self.assertTrue(sheet._images, 'The QA-selected source photo should remain in the existing report')
        workbook.close()
        self.checklist.refresh_from_db(); revision.refresh_from_db()
        self.assertEqual(self.checklist.answer_data, {})
        self.assertEqual(revision.snapshot, immutable_snapshot)

    def test_excluded_template_can_submit_category_evidence_without_enabling_drawings(self):
        excluded = ChecklistTemplate.objects.create(name='SYNTHETIC excluded category checklist',
            definition=deepcopy(self.schema), published=True, checklist_required=False)
        self.project.checklist_template = excluded
        self.project.save(update_fields=['checklist_template'])
        self.checklist.checklist_template = excluded
        self.checklist.save(update_fields=['checklist_template'])
        WorkflowTemplateQualification.objects.create(project=self.project, template=excluded,
            workflow_type=self.checklist.workflow_type, definition_digest=definition_digest(self.schema),
            criteria_reference='SYNTHETIC excluded survey criteria', inspection_method='Read source labels',
            reviewed_by=self.lead, reason='Keep original drawing scope excluded')
        self.upload('Civil')
        submitted, _ = self.submit()
        self.assertIn(submitted.status_code, (200, 201), submitted.content)
        revision = self.assert_submitted()
        _, inference, renderer = self.processing()
        self.assertTrue(pipeline.process_next())
        revision.processing_job.refresh_from_db()
        self.assertEqual(revision.processing_job.status, 'ready_for_review', revision.processing_job.error)
        self.assertTrue(inference.called)
        renderer.assert_not_called()
        self.assertFalse(revision.drawing_sets.exists())
        excluded.refresh_from_db()
        self.assertFalse(excluded.checklist_required)

    def test_qa_rejects_stale_edits_unknown_fields_and_foreign_sources(self):
        self.upload('Civil')
        submitted, _ = self.submit()
        self.assertIn(submitted.status_code, (200, 201), submitted.content)
        revision = self.assert_submitted()
        original = self.qa_state()
        saved = self.qa_action('answer', field_id='contact', value='Checked contact',
            photo_ids=[], reason='Synthetic recorded contact')
        self.assertEqual(saved.status_code, 200, saved.content)
        events = revision.processing_events.count()
        stale = self.post('photo_capture_review_action', {'action': 'answer',
            'expected_review_digest': original['review_digest'], 'field_id': 'contact',
            'value': 'Stale overwrite', 'reason': 'Stale request', 'photo_ids': []}, actor=self.quality)
        self.assertEqual(stale.status_code, 409, stale.content)
        for field, ids in [('not_a_published_field', []), ('manufacturer', [self.photo.pk])]:
            invalid = self.qa_action('answer', field_id=field, value='Unsupported overwrite',
                photo_ids=ids, reason='Invalid fixture request')
            self.assertEqual(invalid.status_code, 409, invalid.content)
        self.assertEqual(revision.processing_events.count(), events)
        self.assertEqual(next(item for item in self.qa_state()['items'] if item['id'] == 'contact')['value'], 'Checked contact')

    def test_legacy_engineer_http_ai_actions_are_forbidden_with_real_existing_proposals(self):
        legacy = self.photo.checklist
        self.assertNotIn('category_intake_version', legacy.intake_metadata)
        original_answers = deepcopy(legacy.answer_data)
        retained_job = PhotoJob.objects.create(checklist=legacy, requested_by=self.engineer,
            schema=legacy.schema_snapshot, model='synthetic-historical-model', status='failed')
        retained_job.photos.add(self.photo)
        draft = PhotoDraft.objects.create(job=retained_job, photo=self.photo,
            field_id='manufacturer', label='Cabinet manufacturer', value='SyntheticCo',
            reason='Synthetic pre-existing proposal', uncertainty='low', baseline=None)
        snapshot = pipeline.freeze_data(legacy)
        revision = SurveyProcessingRevision.objects.create(checklist=legacy, number=1,
            snapshot=snapshot, digest=pipeline.digest(snapshot), submitted_by=self.engineer)
        processing_job = SurveyProcessingJob.objects.create(revision=revision,
            model='synthetic-retained-model', status='ready_for_review', completed=2, total=2)
        payload = pipeline.prepare_checklist(snapshot, {'findings': [], 'missing': [], 'conflicts': []})
        SurveyProcessingResult.objects.create(revision=revision, payload=payload, digest=pipeline.digest(payload))
        self.client.force_login(self.engineer)
        mapping = self.client.post(reverse('photo_mapping', args=[legacy.pk]), json.dumps({
            'expected_schema_digest': schema_digest(legacy), 'decisions': [{
                'photo_id': self.photo.pk, 'field_id': 'manufacturer', 'action': 'manual'}]}),
            content_type='application/json')
        self.assertEqual(mapping.status_code, 403, mapping.content)
        for route, args, data in [
            ('photo_accept', [legacy.pk, draft.pk], {}),
            ('photo_retry', [legacy.pk, retained_job.pk], {}),
            ('photo_enqueue', [legacy.pk], {'photo_ids': [self.photo.pk]}),
        ]:
            with self.subTest(route=route):
                denied = self.client.post(reverse(route, args=args), json.dumps(data), content_type='application/json')
                self.assertEqual(denied.status_code, 403, denied.content)
        for action, extra in [('review', {'decision': 'accept', 'finding_ids': json.dumps(['manufacturer'])}),
                              ('review', {'decision': 'engineer_complete'}),
                              ('submit', {'uploads_finished': 'true'})]:
            with self.subTest(action=action, decision=extra.get('decision')):
                denied = self.client.post(reverse('survey_processing_action', args=[legacy.pk]),
                    {'action': action, 'revision_id': revision.pk, **extra})
                self.assertEqual(denied.status_code, 403, denied.content)
        processing_job.status = 'failed'
        processing_job.save(update_fields=['status'])
        denied = self.client.post(reverse('survey_processing_action', args=[legacy.pk]),
            {'action': 'retry', 'revision_id': revision.pk})
        self.assertEqual(denied.status_code, 403, denied.content)
        bookmark = self.client.get(reverse('survey_processing_workspace', args=[legacy.pk]))
        self.assertRedirects(bookmark, reverse('photo_workspace', args=[legacy.pk]), fetch_redirect_response=False)
        legacy.refresh_from_db(); draft.refresh_from_db(); retained_job.refresh_from_db(); revision.refresh_from_db()
        self.assertNotIn('category_intake_version', legacy.intake_metadata)
        self.assertEqual(legacy.answer_data, original_answers)
        self.assertEqual(draft.status, 'pending')
        self.assertIsNone(draft.accepted_by_id)
        self.assertEqual(retained_job.status, 'failed')
        self.assertFalse(PhotoPlacement.objects.filter(photo=self.photo).exists())
        self.assertFalse(revision.processing_events.exists())
        self.assertEqual(revision.snapshot, snapshot)
