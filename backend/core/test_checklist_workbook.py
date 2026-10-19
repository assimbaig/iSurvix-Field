"""Parser tests author only in-memory adversarial XLSX fixtures, never artifacts."""
from copy import deepcopy
from io import BytesIO
import json
import zipfile
from unittest.mock import patch
from django.test import SimpleTestCase
from openpyxl import Workbook
from .checklist_workbook import workbook_rows,parse_workbook,WorkbookValidationError,HEADERS,MAX_BYTES,export_workbook
from .standard_checklists import build_standard
from .template_service import default_definition

class ChecklistWorkbookTests(SimpleTestCase):
    def record(self,key='one',definition=None):return {'key':key,'name':'Example checklist','industry':'telecom','workflow_type':'telecom_survey','definition':definition or default_definition()}
    def xlsx(self,records=None,edit=None):
        data=workbook_rows(records or [self.record()]);wb=Workbook();wb.remove(wb.active)
        for name,table in data.items():
            sheet=wb.create_sheet(name);sheet.append(table['headers'])
            for row in table['rows']:sheet.append(row)
        if edit:edit(wb)
        output=BytesIO();wb.save(output);wb.close();return output.getvalue()
    def assert_bad(self,data,sheet=None,column=None):
        with self.assertRaises(WorkbookValidationError) as raised:parse_workbook(data)
        errors=raised.exception.errors;self.assertTrue(errors)
        self.assertTrue(all(set(e)=={'sheet','row','column','message'} for e in errors))
        if sheet:self.assertEqual(errors[0]['sheet'],sheet)
        if column:self.assertEqual(errors[0]['column'],column)
        return errors
    def test_original_master_lossless_and_no_database_access(self):
        record=self.record();before=deepcopy(record);payload=self.xlsx()
        with patch('django.db.backends.utils.CursorWrapper.execute',side_effect=AssertionError('No database access')):
            result=parse_workbook(payload)
        self.assertEqual(result['checklists'],[before]);self.assertEqual(record,before)
        self.assertEqual(len(result['sha256']),64);self.assertEqual(result['warnings'],[])
    def test_all_four_standards_roundtrip_with_room_context(self):
        keys=['telecom_survey_default','telecom_inspection_bts','telecom_inspection_colocation','construction_inspection_all'];records=[]
        for i,key in enumerate(keys):
            definition=build_standard(key,rooms=[{'type':'bathroom','name':'Bath 1','floor':'Ground'}] if key.startswith('construction') else None)
            meta=definition['standard_catalog'];records.append({'key':'standard'+str(i),'name':key,'industry':meta['industry'],'workflow_type':meta['workflow_type'],'definition':definition})
        self.assertEqual(parse_workbook(self.xlsx(records))['checklists'],records)
    def test_supported_field_types_rules_and_empty_descriptive_values_roundtrip(self):
        fields=[{'id':'f'+str(i),'label':kind,'type':kind,'required':False} for i,kind in enumerate(['text','textarea','number','date','select','inspection','table','photo'])]
        fields[0].update(system='site_id',description='',unit='m')
        fields[4]['options']=['One','Two'];fields[6]['columns']=[{'id':'col','label':'Column','unit':'m'}]
        fields[5]['evidence_rules']={'version':1,'required_subjects':['tower'],'any_subjects':['ladder'],'required_text':['SITE'],'forbidden_subjects':['rust'],'expected_site_id':'SITE-1','require_readable_text':True,'require_qr':False,'require_gps':True,'require_capture_time':True,'min_photos':2,'minimum_quality':'clear','notes':'Review original'}
        record=self.record(definition={'title':'Custom','description':'','sections':[{'id':'general','title':'General','group':'','fields':fields}]})
        self.assertEqual(parse_workbook(self.xlsx([record]))['checklists'],[record])
    def test_formulas_errors_unknown_headers_and_unknown_sheets_rejected(self):
        for value in ['=HYPERLINK("https://example.invalid","x")','#DIV/0!']:
            self.assert_bad(self.xlsx(edit=lambda wb: setattr(wb['Items']['D2'],'value',value)),'Items','label')
        self.assert_bad(self.xlsx(edit=lambda wb: setattr(wb['Items']['N1'],'value','skip_if')),'Items')
        self.assert_bad(self.xlsx(edit=lambda wb: wb.create_sheet('HiddenConditions')),'HiddenConditions')
    def test_unknown_extras_and_invalid_rules_not_silently_accepted(self):
        col=HEADERS['Items'].index('extra_json')+1
        self.assert_bad(self.xlsx(edit=lambda wb: setattr(wb['Items'].cell(2,col),'value','{"auto_approve":true}')),'Items','extra_json')
        def edit(wb):
            values={'checklist_key':'one','field_id':wb['Items']['C2'].value,'version':1,'min_photos':True}
            wb['EvidenceRules'].append([values.get(key) for key in HEADERS['EvidenceRules']])
        self.assert_bad(self.xlsx(edit=edit),'EvidenceRules','min_photos')
    def test_duplicate_and_unresolved_references_have_locations(self):
        self.assert_bad(self.xlsx(edit=lambda wb: wb['Items'].append([c.value for c in wb['Items'][2]])),'Items','field_id')
        self.assert_bad(self.xlsx(edit=lambda wb: setattr(wb['Items']['B2'],'value','missing')),'Items','section_id')
        self.assert_bad(self.xlsx(edit=lambda wb: setattr(wb['Sections']['A2'],'value','foreign')),'Sections','checklist_key')
    def test_invalid_choice_orders_and_type_references(self):
        def edit(wb):wb['Choices'].append(['one','site_id','Bad',1])
        self.assert_bad(self.xlsx(edit=edit),'Choices','field_id')
    def test_macro_external_relationship_and_xml_entity_packages_rejected(self):
        source=self.xlsx()
        def add(name,value):
            result=BytesIO()
            with zipfile.ZipFile(BytesIO(source)) as src,zipfile.ZipFile(result,'w',zipfile.ZIP_DEFLATED) as dst:
                for item in src.infolist():dst.writestr(item.filename,src.read(item))
                dst.writestr(name,value)
            return result.getvalue()
        self.assert_bad(add('xl/vbaProject.bin',b'macro'))
        self.assert_bad(add('xl/worksheets/_rels/extra.rels','<Relationships><Relationship TargetMode="External" Target="https://example.invalid" /></Relationships>'))
        self.assert_bad(add('xl/entities.xml','<!DOCTYPE a [<!ENTITY b "expanded">]><a>&b;</a>'))
    def test_size_and_corrupt_package_bounds(self):
        for payload in [b'',b'not a zip',b'x'*(MAX_BYTES+1)]:self.assert_bad(payload)
        with patch('core.checklist_workbook.MAX_EXPANDED',1):self.assert_bad(self.xlsx())
    def test_checklist_and_field_total_bounds(self):
        records=[self.record('c'+str(i)) for i in range(10)]
        self.assertEqual(len(parse_workbook(self.xlsx(records))['checklists']),10)
        with self.assertRaises(WorkbookValidationError):workbook_rows(records+[self.record('eleven')])
        def edit(wb):
            wb['Checklists'].append(['c11','Extra','telecom','telecom_survey'])
        self.assert_bad(self.xlsx(records,edit),'Checklists','checklist_key')
    def test_unknown_source_metadata_cannot_disappear_on_export(self):
        record=self.record();record['definition']['skip_if']='unsafe unsupported condition'
        with self.assertRaises(WorkbookValidationError):workbook_rows([record])
    def test_mutating_export_rows_does_not_mutate_source(self):
        record=self.record();before=deepcopy(record);rows=workbook_rows([record]);rows['Items']['rows'][0][3]='Changed'
        self.assertEqual(record,before)

    def test_runtime_writer_preserves_literal_formula_like_text_and_dropdowns(self):
        from openpyxl import load_workbook
        record=self.record();record['name']='=Literal name'
        record['definition']['sections'][0]['fields'][0]['label']='=Literal label'
        payload=export_workbook([record])
        self.assertEqual(parse_workbook(payload)['checklists'],[record])
        workbook=load_workbook(BytesIO(payload),data_only=False)
        self.assertEqual(workbook['Items']['D2'].data_type,'s')
        self.assertTrue(workbook['Items'].data_validations.count)
        self.assertEqual(workbook['_Lists'].sheet_state,'hidden')
        self.assertIn('choices_field_type',workbook.defined_names)
        workbook.close()

    def test_wizard_context_roundtrip_and_mismatch_denied(self):
        record=self.record();record['definition']['wizard_context']={'version':1,'industry':'telecom','workflow_type':'telecom_survey'}
        self.assertEqual(parse_workbook(export_workbook([record]))['checklists'],[record])
        record['definition']['wizard_context']['industry']='construction'
        with self.assertRaises(WorkbookValidationError):export_workbook([record])

    def test_granular_construction_report_tables_roundtrip(self):
        definition=build_standard('construction_inspection_all',stages=['excavation','field_density_test','foundation_concrete','other_floor_concrete'])
        record={'key':'construction','name':'Construction activities','industry':'construction','workflow_type':'construction_inspection','definition':definition}
        self.assertEqual(parse_workbook(export_workbook([record]))['checklists'],[record])

    def test_missing_or_underreported_optional_dimensions_do_not_drop_data(self):
        import re
        original=self.xlsx()
        def rewritten(dimension):
            output=BytesIO()
            with zipfile.ZipFile(BytesIO(original)) as source,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as target:
                for entry in source.infolist():
                    data=source.read(entry)
                    if entry.filename.startswith('xl/worksheets/') and entry.filename.endswith('.xml'):
                        data=re.sub(rb'<dimension[^>]*/>',dimension,data)
                    target.writestr(entry.filename,data)
            return output.getvalue()
        self.assertEqual(parse_workbook(rewritten(b''))['checklists'],[self.record()])
        self.assertEqual(parse_workbook(rewritten(b'<dimension ref="A1:A1"/>'))['checklists'],[self.record()])
