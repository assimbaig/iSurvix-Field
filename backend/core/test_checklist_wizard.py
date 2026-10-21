"""Isolated wizard HTTP contracts, import security and existing-data preservation."""
import json
import tempfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from datetime import timedelta
from unittest.mock import patch
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client,TestCase,override_settings
from django.urls import reverse
from django.utils import timezone
from .models import Company,Project,Profile,IndustryCatalog,ChecklistTemplate,Checklist
from .template_service import default_definition


@override_settings(SECRET_KEY='isolated-checklist-wizard-tests',ROOT_URLCONF='checklist.urls_local',
                   ALLOWED_HOSTS=['testserver'],SECURE_SSL_REDIRECT=False,
                   PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class ChecklistWizardTests(TestCase):
    def setUp(self):
        folder=Path(__file__).resolve().parents[2]/'.local/tmp';folder.mkdir(parents=True,exist_ok=True)
        temporary=tempfile.TemporaryDirectory(dir=folder,prefix='checklist-wizard-')
        self.addCleanup(temporary.cleanup);storage=override_settings(MEDIA_ROOT=temporary.name)
        storage.enable();self.addCleanup(storage.disable)
        self.company=Company.objects.create(name='Synthetic wizard company')
        self.source=ChecklistTemplate.objects.create(name='Preserved published master',definition=default_definition(),published=True)
        self.project=Project.objects.create(name='Wizard scoped project',company=self.company,industry='telecom',checklist_template=self.source)
        self.admin=self.person('admin',Profile.Roles.ADMIN)
        self.draft=ChecklistTemplate.objects.create(name='Owned existing draft',definition=default_definition(),created_by=self.admin)
        self.other_company=Company.objects.create(name='Foreign wizard company')
        self.other_project=Project.objects.create(name='Foreign wizard project',company=self.other_company,industry='telecom')
        self.other_admin=self.person('foreign_admin',Profile.Roles.ADMIN,self.other_project)
        self.foreign=ChecklistTemplate.objects.create(name='Foreign private draft',definition=default_definition(),created_by=self.other_admin)
        self.superuser=self.person('super',Profile.Roles.SUPER_ADMIN,superuser=True)
        self.engineer=self.person('engineer',Profile.Roles.ENGINEER)
        self.quality=self.person('quality',Profile.Roles.QUALITY_ENGINEER)
        self.lead=self.person('lead',Profile.Roles.TEAM_LEAD)
        self.customer=self.person('customer',Profile.Roles.CLIENT)
        self.survey=Checklist.objects.create(project=self.project,user=self.engineer,quality_reviewer=self.quality,
            site_id='PRESERVED-WIZARD-SITE',answer_data={'custom':{'existing':'Human answer'}},
            section_approvals={'old':{'by':self.lead.pk}},review_history=[{'action':'historical'}])
        self.client.force_login(self.admin)

    def person(self,name,role,project=None,superuser=False):
        project=project or self.project
        actor=User.objects.create_user('wizard_'+name,is_superuser=superuser)
        actor.profile.role=role;actor.profile.project=project;actor.profile.company=project.company
        if role==Profile.Roles.ADMIN:actor.profile.admin_industry=IndustryCatalog.objects.get(code=project.industry)
        actor.profile.save();return actor

    def post(self,action,**data):
        return self.client.post(reverse('checklist_wizard'),{'action':action,**data})

    def before(self):
        self.source.refresh_from_db();self.survey.refresh_from_db();self.project.refresh_from_db()
        return deepcopy((self.source.definition,self.source.published,self.source.family_id,
                         self.survey.schema_snapshot,self.survey.answer_data,self.survey.section_approvals,
                         self.survey.review_history,self.project.checklist_template_id))

    def records(self):
        return [{'key':'survey_one','name':'Workbook first draft','industry':'telecom','workflow_type':'telecom_survey',
                 'definition':{'sections':[{'id':'general','title':'First section','fields':[
                     {'id':'condition','label':'Visible foundation condition','type':'inspection','required':True,'photo_required':True,
                      'evidence_rules':{'version':1,'required_subjects':['concrete foundation'],'require_gps':True,'min_photos':1}}]}]}},
                {'key':'survey_two','name':'Workbook second draft','industry':'telecom','workflow_type':'telecom_inspection',
                 'definition':{'sections':[{'id':'general','title':'Second section','fields':[
                     {'id':'equipment','label':'Equipment type','type':'select','options':['Cabinet','Antenna'],'required':False}]}]}}]

    def workbook(self,records=None,mutate=None):
        from openpyxl import Workbook
        from .checklist_workbook import workbook_rows
        workbook=Workbook();workbook.remove(workbook.active)
        for sheet,table in workbook_rows(records or self.records()).items():
            page=workbook.create_sheet(sheet);page.append(table['headers'])
            for row in table['rows']:page.append(row)
        if mutate:mutate(workbook)
        data=BytesIO();workbook.save(data);return data.getvalue()

    def upload(self,raw=None,filename='synthetic-multi-checklists.xlsx'):
        return SimpleUploadedFile(filename,raw if raw is not None else self.workbook(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def stage(self,raw=None):
        response=self.post('preview_import',workbook=self.upload(raw))
        self.assertEqual(response.status_code,200,response.content)
        self.assertIsNotNone(response.context['import_batch'])
        return response.context['import_batch']

    def commit(self,batch,digest=None):
        return self.post('commit_import',import_id=str(batch.pk),expected_digest=digest or batch.payload_digest)

    def test_super_can_read_landing_and_download_workbook_but_cannot_run_engineering_actions(self):
        self.client.force_login(self.superuser)
        self.assertEqual(self.client.get(reverse('checklist_wizard')).status_code,200)
        response=self.client.get(reverse('checklist_wizard_workbook'))
        self.assertEqual(response.status_code,200)
        self.assertIn('spreadsheetml',response['Content-Type'])
        before=ChecklistTemplate.objects.count()
        for action,data in [('create_blank',{'name':'Denied','industry':'telecom','workflow_type':'telecom_survey'}),
                            ('replicate',{'source_id':self.source.pk,'name':'Denied'}),
                            ('preview_import',{}),('commit_import',{'import_id':999,'expected_digest':'forged'})]:
            self.assertEqual(self.post(action,**data).status_code,403,action)
        self.assertEqual(self.client.get(reverse('checklist_wizard_edit',args=[self.draft.pk])).status_code,403)
        self.assertEqual(self.client.get(reverse('checklist_wizard_export',args=[self.source.pk])).status_code,403)
        self.assertEqual(ChecklistTemplate.objects.count(),before)

    def test_other_roles_locked_and_forged_super_profile_are_denied(self):
        for actor in [self.engineer,self.quality,self.lead,self.customer]:
            self.client.force_login(actor)
            self.assertEqual(self.client.get(reverse('checklist_wizard')).status_code,403)
            self.assertEqual(self.post('create_blank',name='Denied',industry='telecom',workflow_type='telecom_survey').status_code,403)
        self.admin.profile.is_locked=True;self.admin.profile.save();self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse('checklist_wizard')).status_code,403)
        self.engineer.profile.role=Profile.Roles.SUPER_ADMIN;self.engineer.profile.save();self.client.force_login(self.engineer)
        self.assertEqual(self.client.get(reverse('checklist_wizard')).status_code,403)

    def test_blank_creation_is_scoped_new_draft_and_never_changes_existing_records(self):
        before=self.before()
        response=self.post('create_blank',name='Blank synthetic wizard',industry='telecom',workflow_type='telecom_survey')
        self.assertEqual(response.status_code,302,response.content)
        draft=ChecklistTemplate.objects.get(name='Blank synthetic wizard')
        self.assertFalse(draft.published);self.assertEqual(draft.version,1);self.assertEqual(draft.created_by,self.admin)
        self.assertNotEqual(draft.family_id,self.source.family_id)
        self.assertEqual(self.before(),before)
        self.assertIn(self.post('create_blank',name='Wrong industry',industry='construction',workflow_type='construction_inspection').status_code,[400,403])
        self.assertFalse(ChecklistTemplate.objects.filter(name='Wrong industry').exists())

    def test_copy_and_edit_are_scoped_and_published_source_remains_immutable(self):
        before=self.before()
        response=self.post('replicate',source_id=self.source.pk,name='Independent source copy')
        self.assertEqual(response.status_code,302,response.content)
        copied=ChecklistTemplate.objects.get(name='Independent source copy')
        self.assertFalse(copied.published);self.assertEqual(copied.version,1)
        self.assertNotEqual(copied.family_id,self.source.family_id)
        self.assertEqual(copied.definition['sections'],self.source.definition['sections'])
        self.assertIn(self.post('replicate',source_id=self.foreign.pk,name='Foreign copy').status_code,[403,404])
        self.assertEqual(self.client.get(reverse('checklist_wizard_edit',args=[self.draft.pk])).status_code,200)
        published=self.client.get(reverse('checklist_wizard_edit',args=[self.source.pk]))
        self.assertEqual(published.status_code,200)
        refused=self.client.post(reverse('checklist_wizard_edit',args=[self.source.pk]),
            {'action':'save','name':'Forbidden published overwrite','definition':json.dumps(self.source.definition),
             'wizard_expected_digest':published.context['wizard_expected_digest']})
        self.assertEqual(refused.status_code,400)
        self.assertFalse(ChecklistTemplate.objects.filter(name='Forbidden published overwrite').exists())
        self.assertIn(self.client.get(reverse('checklist_wizard_edit',args=[self.foreign.pk])).status_code,[403,404])
        self.assertEqual(self.before(),before)

    def test_multi_preview_has_no_side_effect_and_commit_is_atomic_and_replay_idempotent(self):
        before=self.before();count=ChecklistTemplate.objects.count()
        batch=self.stage()
        self.assertEqual(ChecklistTemplate.objects.count(),count)
        self.assertEqual(len(batch.payload['checklists']),2)
        self.assertIn(self.commit(batch,'forged-digest').status_code,[400,409])
        self.assertEqual(ChecklistTemplate.objects.count(),count)
        response=self.commit(batch);self.assertEqual(response.status_code,302,response.content)
        batch.refresh_from_db();first_ids=list(batch.created_template_ids)
        self.assertEqual(len(first_ids),2)
        self.assertEqual(ChecklistTemplate.objects.count(),count+2)
        self.assertEqual(self.commit(batch).status_code,302)
        batch.refresh_from_db();self.assertEqual(batch.created_template_ids,first_ids)
        self.assertEqual(ChecklistTemplate.objects.count(),count+2)
        templates=list(ChecklistTemplate.objects.filter(pk__in=first_ids))
        self.assertTrue(all(not template.published and template.version==1 and template.created_by_id==self.admin.pk for template in templates))
        imported=next(t for t in templates if t.name=='Workbook first draft')
        field=imported.definition['sections'][0]['fields'][0]
        self.assertEqual(field['id'],'condition')
        self.assertEqual(field['evidence_rules']['required_subjects'],['concrete foundation'])
        self.assertTrue(field['evidence_rules']['require_gps'])
        self.assertEqual(self.before(),before)

    def test_two_draft_commit_rolls_back_every_created_template_on_second_save_failure(self):
        batch=self.stage();count=ChecklistTemplate.objects.count();save=ChecklistTemplate.save
        def fail_second(instance,*args,**kwargs):
            if instance.name=='Workbook second draft':raise ValidationError('Synthetic second checklist validation failure')
            return save(instance,*args,**kwargs)
        with patch.object(ChecklistTemplate,'save',fail_second):
            response=self.commit(batch)
        self.assertIn(response.status_code,[400,409],response.content)
        self.assertEqual(ChecklistTemplate.objects.count(),count)
        self.assertFalse(ChecklistTemplate.objects.filter(name='Workbook first draft').exists())
        batch.refresh_from_db();self.assertEqual(batch.created_template_ids,[])
        self.assertIsNone(batch.committed_at)

    def test_staged_import_owner_expiry_company_and_current_account_controls_are_rechecked(self):
        batch=self.stage();count=ChecklistTemplate.objects.count()
        self.client.force_login(self.other_admin)
        self.assertIn(self.commit(batch).status_code,[403,404])
        self.client.force_login(self.admin)
        with patch('django.utils.timezone.now',return_value=batch.expires_at+timedelta(seconds=1)):
            self.assertIn(self.commit(batch).status_code,[400,403,409])
        self.admin.profile.is_locked=True;self.admin.profile.save()
        self.assertEqual(self.commit(batch).status_code,403)
        self.admin.profile.is_locked=False;self.admin.profile.company=self.other_company;self.admin.profile.save()
        self.assertIn(self.commit(batch).status_code,[403,404])
        self.admin.profile.company=self.company;self.admin.profile.save()
        from .models import ProjectMembership
        ProjectMembership.objects.create(user=self.admin,project=self.project,is_active=False)
        self.assertEqual(self.commit(batch).status_code,403)
        self.assertEqual(ChecklistTemplate.objects.count(),count)

    def test_invalid_row_formula_and_approval_columns_report_real_errors_without_partial_import(self):
        from .models import ChecklistImportBatch
        before=ChecklistTemplate.objects.count();staged=ChecklistImportBatch.objects.count()
        def duplicate(workbook):
            sheet=workbook['Items'];sheet.append([cell.value for cell in sheet[2]])
        def formula(workbook):workbook['Items']['D2']='=1+1'
        def approval_column(workbook):
            sheet=workbook['Items'];sheet.cell(1,sheet.max_column+1,'approval_status')
        for mutate in [duplicate,formula,approval_column]:
            response=self.post('preview_import',workbook=self.upload(self.workbook(mutate=mutate)))
            self.assertEqual(response.status_code,400,response.content)
            self.assertContains(response,'Items',status_code=400)
            self.assertTrue(response.context['import_errors'])
        self.assertEqual(ChecklistImportBatch.objects.count(),staged)
        self.assertEqual(ChecklistTemplate.objects.count(),before)

    def test_mixed_foreign_industry_workbook_is_never_partially_created(self):
        records=self.records();records[1]['industry']='construction';records[1]['workflow_type']='construction_inspection'
        count=ChecklistTemplate.objects.count()
        response=self.post('preview_import',workbook=self.upload(self.workbook(records)))
        self.assertIn(response.status_code,[400,403],response.content)
        self.assertEqual(ChecklistTemplate.objects.count(),count)

    def test_exported_master_roundtrips_all_fields_into_new_draft_and_source_is_unchanged(self):
        before=self.before()
        response=self.client.get(reverse('checklist_wizard_export',args=[self.source.pk]))
        self.assertEqual(response.status_code,200,response.content if not response.streaming else '')
        raw=b''.join(response.streaming_content) if response.streaming else response.content
        batch=self.stage(raw);self.assertEqual(self.commit(batch).status_code,302)
        batch.refresh_from_db();self.assertEqual(len(batch.created_template_ids),1)
        imported=ChecklistTemplate.objects.get(pk=batch.created_template_ids[0])
        self.assertEqual(imported.definition['sections'],self.source.definition['sections'])
        self.assertEqual(sum(len(s['fields']) for s in imported.definition['sections']),
                         sum(len(s['fields']) for s in default_definition()['sections']))
        self.assertEqual(self.before(),before)

    def test_csrf_positive_file_preview_and_commit_work_but_missing_token_is_denied(self):
        secure=Client(enforce_csrf_checks=True);secure.force_login(self.admin)
        secure.get(reverse('checklist_wizard'));token=secure.cookies['csrftoken'].value
        self.assertEqual(secure.post(reverse('checklist_wizard'),{'action':'preview_import','workbook':self.upload()}).status_code,403)
        response=secure.post(reverse('checklist_wizard'),{'action':'preview_import','workbook':self.upload(),'csrfmiddlewaretoken':token})
        self.assertEqual(response.status_code,200,response.content)
        batch=response.context['import_batch']
        self.assertEqual(secure.post(reverse('checklist_wizard'),{'action':'commit_import','import_id':batch.pk,'expected_digest':batch.payload_digest}).status_code,403)
        response=secure.post(reverse('checklist_wizard'),{'action':'commit_import','import_id':batch.pk,
            'expected_digest':batch.payload_digest,'csrfmiddlewaretoken':token})
        self.assertEqual(response.status_code,302,response.content)

    def test_preview_input_immutability_and_digest_check_reject_tampering_without_overwrite(self):
        from .models import ChecklistImportBatch
        batch=self.stage();original=deepcopy(batch.payload);count=ChecklistTemplate.objects.count()
        batch.payload={'checklists':[]}
        with self.assertRaises(ValidationError):batch.save()
        batch.refresh_from_db();self.assertEqual(batch.payload,original)
        # Simulate storage corruption only in the disposable test database.
        ChecklistImportBatch.objects.filter(pk=batch.pk).update(payload={'checklists':[]})
        self.assertIn(self.commit(batch).status_code,[400,409])
        self.assertEqual(ChecklistTemplate.objects.count(),count)
        self.client.force_login(self.other_admin)
        self.assertIn(self.client.get(reverse('checklist_wizard'),{'import_id':str(batch.pk)}).status_code,[403,404])

    def test_wizard_editor_stale_digest_preserves_first_save_and_retains_unsaved_second_payload(self):
        url=reverse('checklist_wizard_edit',args=[self.draft.pk])
        page=self.client.get(url);self.assertEqual(page.status_code,200)
        token=page.context['wizard_expected_digest'];first=deepcopy(self.draft.definition)
        first['sections'][0]['fields'][0]['label']='First editor source label'
        saved=self.client.post(url,{'action':'save','name':'First editor wins','definition':json.dumps(first),'wizard_expected_digest':token})
        self.assertEqual(saved.status_code,302,saved.content)
        second=deepcopy(first);second['sections'][0]['fields'][0]['label']='Second unsaved source label'
        rejected=self.client.post(url,{'action':'save','name':'Second stale editor','definition':json.dumps(second),'wizard_expected_digest':token})
        self.assertEqual(rejected.status_code,409,rejected.content)
        self.draft.refresh_from_db();self.assertEqual(self.draft.name,'First editor wins')
        self.assertEqual(self.draft.definition['sections'][0]['fields'][0]['label'],'First editor source label')
        self.assertEqual(rejected.context['definition']['sections'][0]['fields'][0]['label'],'Second unsaved source label')

    def test_documented_default_download_is_genuinely_importable_multiple_checklists(self):
        from .checklist_workbook import parse_workbook
        before=self.before()
        response=self.client.get(reverse('checklist_wizard_workbook'));self.assertEqual(response.status_code,200)
        raw=b''.join(response.streaming_content) if response.streaming else response.content
        parsed=parse_workbook(raw);self.assertGreaterEqual(len(parsed['checklists']),2)
        batch=self.stage(raw);self.assertEqual(self.commit(batch).status_code,302)
        batch.refresh_from_db();self.assertEqual(len(batch.created_template_ids),len(parsed['checklists']))
        self.assertEqual(self.before(),before)

    def test_macros_oversized_upload_and_unknown_conditions_are_blocked_at_http_boundary(self):
        import zipfile
        from .models import ChecklistImportBatch
        count=ChecklistTemplate.objects.count();staged=ChecklistImportBatch.objects.count()
        macro=BytesIO(self.workbook())
        with zipfile.ZipFile(macro,'a') as archive:archive.writestr('xl/vbaProject.bin',b'Synthetic blocked macro marker')
        def unknown_condition(workbook):
            sheet=workbook['EvidenceRules'];sheet.cell(1,sheet.max_column+1,'when_other_answer_equals')
        for raw in [macro.getvalue(),b'0'*(5*1024*1024+1),self.workbook(mutate=unknown_condition)]:
            response=self.post('preview_import',workbook=self.upload(raw))
            self.assertEqual(response.status_code,400,response.content)
        self.assertEqual(ChecklistImportBatch.objects.count(),staged)
        self.assertEqual(ChecklistTemplate.objects.count(),count)

    def test_legacy_editor_cannot_change_wizard_industry_or_remove_standard_identity(self):
        from .standard_checklists import build_standard
        original=deepcopy(self.draft.definition)
        original['wizard_context']={'version':1,'industry':'telecom','workflow_type':'telecom_survey'}
        self.draft.definition=original;self.draft.save()
        forged=deepcopy(original);forged['wizard_context']['industry']='construction'
        response=self.client.post(reverse('checklist_template_edit',args=[self.draft.pk]),
            {'action':'save','name':'Forged industry','definition':json.dumps(forged)})
        self.assertNotEqual(response.status_code,302)
        self.draft.refresh_from_db();self.assertEqual(self.draft.definition,original)
        standard=ChecklistTemplate.objects.create(name='Owned standard identity',created_by=self.admin,
            definition=build_standard('telecom_inspection_bts'))
        before=deepcopy(standard.definition);removed=deepcopy(before);removed.pop('standard_catalog')
        response=self.client.post(reverse('checklist_template_edit',args=[standard.pk]),
            {'action':'save','name':'Removed source identity','definition':json.dumps(removed)})
        self.assertNotEqual(response.status_code,302)
        standard.refresh_from_db();self.assertEqual(standard.definition,before)

    def test_malformed_import_digest_and_giant_source_id_return_controlled_errors(self):
        batch=self.stage();count=ChecklistTemplate.objects.count()
        for digest in ['not-a-digest','\u00e9'*64,'0'*65]:
            response=self.commit(batch,digest)
            self.assertIn(response.status_code,[400,403,409],response.content)
        for action in ['replicate','edit']:
            response=self.post(action,source_id='9'*5000,name='Rejected giant source')
            self.assertIn(response.status_code,[400,403,404],response.content)
        batch.refresh_from_db();self.assertIsNone(batch.committed_at)
        self.assertEqual(ChecklistTemplate.objects.count(),count)
