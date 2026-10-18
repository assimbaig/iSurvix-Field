import json
import tempfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.template.loader import get_template
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image

from .models import Checklist, ChecklistTemplate, Profile, Project, WorkAssignment
from .template_service import default_definition, duplicate_revision, missing_required, validate_definition


class ChecklistBuilderTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.media = tempfile.TemporaryDirectory()
        cls.override = override_settings(MEDIA_ROOT=cls.media.name)
        cls.override.enable()

    @classmethod
    def tearDownClass(cls):
        cls.override.disable()
        cls.media.cleanup()
        super().tearDownClass()

    def setUp(self):
        self.admin = self.person('fixture_admin', Profile.Roles.ADMIN, 'ADFixture')
        # Operational templates belong to the scoped Project Admin.
        self.engineer = self.person('fixture_engineer', Profile.Roles.ENGINEER, 'EngFixture')
        self.lead = self.person('fixture_lead', Profile.Roles.TEAM_LEAD, 'TLFixture')
        self.quality = self.person('fixture_quality', Profile.Roles.QUALITY_ENGINEER, 'QEFixture')
        self.project = Project.objects.create(name='Fixture project')
        for user in [self.admin, self.engineer, self.lead, self.quality]:
            user.profile.project = self.project
            user.profile.save()
        self.definition = {'title': 'Fixture custom checklist', 'sections': [{'id': 'general', 'title': 'Site details', 'fields': [
            {'id': 'site_id', 'label': 'Site ID', 'type': 'text', 'system': 'site_id', 'required': True},
            {'id': 'height', 'label': 'Tower height', 'type': 'number', 'required': True, 'unit': 'm'},
            {'id': 'condition', 'label': 'Condition', 'type': 'select', 'options': ['Good', 'Repair needed'], 'required': False},
            {'id': 'inspection', 'label': 'Foundation inspected', 'type': 'inspection', 'required': False},
            {'id': 'equipment', 'label': 'Equipment', 'type': 'table', 'required': False, 'columns': [{'id': 'model', 'label': 'Model'}]},
            {'id': 'photo', 'label': 'Site overview', 'type': 'photo', 'required': False},
        ]}]}
        self.template = ChecklistTemplate.objects.create(name='Fixture template', definition=deepcopy(self.definition), created_by=self.admin)
        self.a = Client()
        self.a.force_login(self.admin)
        self.e = Client()
        self.e.force_login(self.engineer)
        self.l = Client()
        self.l.force_login(self.lead)

    def person(self, name, role, path):
        u = User.objects.create_user(username=name, password='Fixture-pass-763!')
        u.profile.role, u.profile.path = role, path
        u.profile.save()
        return u

    def publish_and_create(self):
        self.template.published = True
        self.template.save()
        self.project.checklist_template = self.template
        self.project.save()
        return Checklist.objects.create(project=self.project, user=self.engineer, quality_reviewer=self.quality, site_id='FIXTURE-001')

    def post_answers(self, checklist, answers, client=None):
        return (client or self.e).post(reverse('custom_checklist_save', args=[checklist.id]), json.dumps({'answers': answers}), content_type='application/json')

    def logo(self, name='client.png', color='blue'):
        buffer = BytesIO()
        Image.new('RGB', (100, 30), color).save(buffer, format='PNG')
        return SimpleUploadedFile(name, buffer.getvalue(), content_type='image/png')

    def test_original_172_rows_and_inventory_retained(self):
        schema = default_definition()
        validate_definition(schema)
        fields = [f for s in schema['sections'] for f in s['fields']]
        self.assertEqual(sum('excel_row' in f for f in fields), 172)
        self.assertEqual(sum(f['type'] == 'table' for f in fields), 9)
        self.assertEqual(sum(f['type'] == 'inspection' for f in fields), 150)

    def test_admin_only_template_management(self):
        url = reverse('checklist_templates')
        self.assertEqual(self.a.get(url).status_code, 200)
        self.assertEqual(self.e.get(url).status_code, 403)
        self.assertEqual(self.l.get(url).status_code, 403)
        self.assertEqual(Client().get(url).status_code, 403)

    def test_original_copy_populates_real_client_logos(self):
        response = self.a.post(reverse('checklist_templates'), {'action': 'create_original'})
        self.assertEqual(response.status_code, 302)
        t = ChecklistTemplate.objects.latest('id')
        self.assertTrue(t.client_left_logo.name)
        self.assertTrue(t.client_right_logo.name)
        self.assertEqual(t.client_left_name, 'Nokia')
        self.assertEqual(t.client_right_name, 'TAWAL')

    def test_add_remove_sections_fields_and_replace_logos(self):
        definition = deepcopy(self.definition)
        definition['sections'][0]['fields'].pop(2)
        definition['sections'].append({'id': 'custom', 'title': 'Custom inspection', 'fields': [{'id': 'new_question', 'label': 'Added question', 'type': 'text', 'required': False}]})
        response = self.a.post(reverse('checklist_template_edit', args=[self.template.id]), {'action': 'save', 'name': 'Custom client checklist', 'definition': json.dumps(definition), 'client_left_name': 'Client A', 'client_right_name': 'Client B', 'client_left_logo': self.logo()})
        self.assertEqual(response.status_code, 302)
        self.template.refresh_from_db()
        self.assertEqual(self.template.definition['sections'][1]['fields'][0]['label'], 'Added question')
        self.assertNotIn('condition', [f['id'] for f in self.template.definition['sections'][0]['fields']])
        self.assertEqual(self.template.client_left_name, 'Client A')
        self.assertTrue(self.template.client_left_logo.name)

    def test_invalid_image_is_rejected(self):
        response = self.a.post(reverse('checklist_template_edit', args=[self.template.id]), {'action': 'save', 'name': self.template.name, 'definition': json.dumps(self.definition), 'client_left_logo': SimpleUploadedFile('fake.png', b'not an image', content_type='image/png')})
        self.assertEqual(response.status_code, 200)
        self.template.refresh_from_db()
        self.assertFalse(self.template.client_left_logo.name)

    def test_draft_cannot_be_assigned_or_used(self):
        self.project.checklist_template = self.template
        with self.assertRaises(ValidationError):
            self.project.save()

    def test_published_revisions_locked_and_existing_survey_frozen(self):
        c = self.publish_and_create()
        old = deepcopy(c.schema_snapshot)
        branding = deepcopy(c.branding_snapshot)
        self.template.definition['sections'][0]['title'] = 'Forbidden edit'
        with self.assertRaises(ValidationError):
            self.template.save()
        self.template.refresh_from_db()
        new = duplicate_revision(self.template, self.admin)
        self.assertEqual(new.version, 2)
        new.definition['sections'][0]['title'] = 'Changed future section'
        new.client_left_name = 'Different client'
        new.published = True
        new.save()
        self.project.checklist_template = new
        self.project.save()
        c.refresh_from_db()
        self.assertEqual(c.schema_snapshot, old)
        self.assertEqual(c.branding_snapshot, branding)
        fresh = Checklist.objects.create(project=self.project, user=self.engineer, site_id='FIXTURE-002')
        self.assertEqual(fresh.schema_snapshot['sections'][0]['title'], 'Changed future section')
        self.assertEqual(fresh.branding_snapshot['clients'][0]['name'], 'Different client')

    def test_all_templates_compile_and_custom_form_renders(self):
        from django.conf import settings
        for p in (Path(settings.BASE_DIR) / 'templates').rglob('*.html'):
            get_template(str(p.relative_to(Path(settings.BASE_DIR) / 'templates')))
        c = self.publish_and_create()
        self.assertContains(self.e.get(reverse('checklist_detail', args=[c.id])), 'Fixture custom checklist')
        self.assertContains(self.a.get(reverse('checklist_template_edit', args=[self.template.id])), 'Client logos')

    def test_custom_apis_require_auth_and_project_access(self):
        c = self.publish_and_create()
        self.assertEqual(Client().get(reverse('custom_checklist', args=[c.id])).status_code, 404)
        self.assertEqual(Client().get(reverse('checklist_data_api', args=[c.id])).status_code, 404)
        self.assertEqual(Client().post(reverse('checklist_autosave_api', args=[c.id]), json.dumps({'answers': {'height': '10'}}), content_type='application/json').status_code, 404)
        other = self.person('other_fixture', Profile.Roles.TEAM_LEAD, 'TLOther')
        foreign = Client()
        foreign.force_login(other)
        self.assertEqual(foreign.get(reverse('custom_checklist', args=[c.id])).status_code, 404)
        # A stale session flag cannot grant a non-admin access to a custom survey.
        session = foreign.session
        session['is_dev_admin'] = True
        session.save()
        self.assertEqual(foreign.get(reverse('custom_checklist', args=[c.id])).status_code, 404)

    def test_fields_validated_and_rapid_distinct_saves_retained(self):
        c = self.publish_and_create()
        self.assertEqual(self.post_answers(c, {'height': '40'}).status_code, 200)
        self.assertEqual(self.post_answers(c, {'condition': 'Good'}).status_code, 200)
        c.refresh_from_db()
        self.assertEqual(c.answer_data['custom']['height'], '40')
        self.assertEqual(c.answer_data['custom']['condition'], 'Good')
        for answers in [{'height': 'not a number'}, {'height': 'NaN'}, {'condition': 'Unconfigured choice'}, {'unknown': 'answer'}, {'site_id': 'MISMATCH'}]:
            self.assertEqual(self.post_answers(c, answers).status_code, 400)

    def test_inspection_na_needs_reason(self):
        c = self.publish_and_create()
        self.assertEqual(self.post_answers(c, {'inspection': {'result': 'Not applicable', 'remarks': ''}}).status_code, 400)
        self.assertEqual(self.post_answers(c, {'inspection': {'result': 'Not applicable', 'remarks': 'Asset not installed'}}).status_code, 200)

    def test_repeating_equipment_rows_and_deletion_persist(self):
        c = self.publish_and_create()
        self.assertEqual(self.post_answers(c, {'equipment': [{'model': 'A'}, {'model': 'B'}]}).status_code, 200)
        self.assertEqual(self.post_answers(c, {'equipment': [{'model': 'B'}]}).status_code, 200)
        c.refresh_from_db()
        self.assertEqual(c.answer_data['custom']['equipment'], [{'model': 'B'}])

    def test_submit_review_and_work_status_are_preserved(self):
        c = self.publish_and_create()
        work = WorkAssignment.objects.create(site_id=c.site_id, latitude=24, longitude=46, description='Fixture', assigned_to=self.engineer, assigned_by=self.lead, project=self.project, checklist=c)
        self.e.post(reverse('custom_checklist_submit', args=[c.id]))
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.DRAFT)
        self.post_answers(c, {'height': '40'})
        self.e.post(reverse('custom_checklist_submit', args=[c.id]))
        c.refresh_from_db(); work.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.SUBMITTED)
        self.assertEqual(work.status, WorkAssignment.Status.SUBMITTED)
        from .quality_workflow import review_quality
        review_quality(self.quality,c,'approve',expected_updated_at=c.updated_at.isoformat())
        self.l.post(reverse('checklist_review_update', args=[c.id]), {'status': 'FINAL', 'comment': 'Accepted'})
        c.refresh_from_db(); work.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.FINAL)
        self.assertEqual(work.status, WorkAssignment.Status.COMPLETED)
        self.assertEqual(c.comment_by, self.lead)
        self.assertEqual(self.post_answers(c, {'height': '99'}).status_code, 400)
        self.e.post(reverse('checklist_submit', args=[c.id]))
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.FINAL)

    def test_export_uses_latest_answers_and_current_client_names(self):
        c = self.publish_and_create()
        self.post_answers(c, {'height': '40'})
        first = self.e.get(reverse('custom_checklist_download', args=[c.id]))
        self.assertEqual(first.status_code, 200)
        self.post_answers(c, {'height': '41'})
        second = self.e.get(reverse('engineer_checklist_download', args=['EngFixture', c.id]))
        w = load_workbook(BytesIO(second.content))
        values = [str(cell.value) for row in w.active for cell in row if cell.value is not None]
        self.assertIn('41', values)
        self.assertNotIn('40', values)
        self.assertIn('Nokia', values)
        self.assertIn('TAWAL', values)

    def test_logo_and_evidence_authorization_and_final_lock(self):
        self.template.client_left_logo.save('fixture.png', self.logo(), save=True)
        c = self.publish_and_create()
        self.assertEqual(self.e.get(reverse('custom_client_logo', args=[c.id, 0])).status_code, 200)
        self.assertEqual(Client().get(reverse('custom_client_logo', args=[c.id, 0])).status_code, 404)
        upload_url = reverse('custom_checklist_upload', args=[c.id, 'photo'])
        self.assertEqual(self.e.post(upload_url, {'images': self.logo('evidence.png')}).status_code, 200)
        c.refresh_from_db()
        self.assertEqual(len(c.image_data['photo']), 1)
        photo_url = reverse('custom_checklist_photo', args=[c.id, 'photo', 0])
        self.assertEqual(Client().get(photo_url).status_code, 404)
        self.assertEqual(self.e.get(photo_url).status_code, 200)
        c.status = Checklist.Status.FINAL
        c.save()
        self.assertEqual(self.e.post(upload_url, {'images': self.logo('later.png')}).status_code, 400)
        self.assertEqual(self.e.post(reverse('custom_checklist_photo_delete', args=[c.id, 'photo', 0])).status_code, 400)

    def test_new_assignment_automatically_gets_project_template(self):
        self.publish_and_create()
        response = self.l.post(reverse('assign_work'), {'site_id': 'FIXTURE-NEW', 'latitude': '24.1', 'longitude': '46.1', 'engineer': self.engineer.id, 'project': self.project.id, 'quality_reviewer':self.quality.pk,'description':'Synthetic assigned activity'})
        self.assertEqual(response.status_code, 302)
        c = Checklist.objects.get(site_id='FIXTURE-NEW')
        self.assertEqual(c.checklist_template, self.template)
        self.assertTrue(c.schema_snapshot)

    def test_invalid_schema_cannot_publish(self):
        schema = deepcopy(self.definition)
        schema['sections'][0]['fields'].append(deepcopy(schema['sections'][0]['fields'][0]))
        with self.assertRaises(ValidationError):
            validate_definition(schema)

    def test_manual_engineer_creation_retired_for_assigned_task_workflow(self):
        self.publish_and_create()
        url = reverse('engineer_checklist_new', args=['EngFixture'])
        before = Checklist.objects.count()
        self.assertEqual(self.e.get(url).status_code, 410)
        self.assertEqual(self.e.post(url, {'site_id': 'FIXTURE-MANUAL'}).status_code, 410)
        self.assertEqual(Checklist.objects.count(), before)

    def test_legacy_record_and_project_remain_usable(self):
        legacy_project = Project.objects.create(name='Legacy fixture')
        legacy = Checklist.objects.create(project=legacy_project, user=self.engineer, site_id='LEGACY')
        self.assertEqual(legacy.schema_snapshot, {})
        self.project.checklist_template = None
        self.project.save()
        self.assertEqual(Checklist.objects.get(pk=legacy.pk).status, Checklist.Status.DRAFT)

    def review_fixture(self):
        self.template.definition['sections'].append({'id': 'power', 'title': 'Power supply', 'fields': [{'id': 'volts', 'label': 'Voltage', 'type': 'number', 'required': True}]})
        self.template.save()
        c = self.publish_and_create()
        self.post_answers(c, {'height': '40', 'volts': '48'})
        c.refresh_from_db()
        c.status = Checklist.Status.SUBMITTED
        c.quality_reviewer=self.quality
        c.save(update_fields=['status','quality_reviewer'])
        from .quality_workflow import review_quality
        review_quality(self.quality,c,'approve',expected_updated_at=c.updated_at.isoformat())
        c.refresh_from_db()
        return c

    def review_decision(self, c, action, sections=None, comment='', client=None):
        return (client or self.l).post(reverse('custom_checklist_review', args=[c.pk]), {'action': action, 'sections': sections or [], 'comment': comment})

    def test_partial_review_locks_answers_and_evidence(self):
        c = self.review_fixture()
        self.review_decision(c, 'partial', ['general'])
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.PARTIAL)
        self.assertEqual(list(c.section_approvals), ['general'])
        self.assertEqual(self.post_answers(c, {'height': '55'}).status_code, 400)
        self.assertEqual(self.post_answers(c, {'volts': '49'}).status_code, 200)
        self.assertEqual(self.e.post(reverse('custom_checklist_upload', args=[c.pk, 'photo']), {'images': self.logo()}).status_code, 400)
        self.assertEqual(self.e.post(reverse('custom_checklist_photo_delete', args=[c.pk, 'photo', 0])).status_code, 400)
        page = self.e.get(reverse('custom_checklist', args=[c.pk]))
        self.assertContains(page, 'Checklist section photos')
        self.assertContains(page, 'data-locked="true"')
        self.assertEqual(c.review_history[1]['sections'], ['general'])

    def test_multiple_sections_approval_finishes_assignment(self):
        c = self.review_fixture()
        self.review_decision(c, 'partial', ['general', 'power'])
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.FINAL)
        self.assertEqual(len(c.section_approvals), 2)
        self.assertEqual(self.post_answers(c, {'volts': '50'}).status_code, 400)
        self.review_decision(c, 'reject', comment='Cannot reopen approved record')
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.FINAL)
        self.assertEqual(len(c.review_history), 2)

    def test_full_reject_requires_reason_and_unlocks_partial_sections(self):
        c = self.review_fixture()
        self.review_decision(c, 'partial', ['general'])
        self.review_decision(c, 'reject')
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.PARTIAL)
        self.review_decision(c, 'reject', comment='Missing site evidence')
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.REJECTED)
        self.assertEqual(c.section_approvals, {})
        self.assertEqual(len(c.review_history), 3)
        self.assertEqual(self.post_answers(c, {'height': '45'}).status_code, 200)

    def test_review_permissions_and_invalid_section_selection(self):
        c = self.review_fixture()
        self.assertEqual(self.review_decision(c, 'approve', client=self.e).status_code, 403)
        self.assertEqual(self.review_decision(c, 'approve', client=Client()).status_code, 404)
        other = self.person('other_lead', Profile.Roles.TEAM_LEAD, 'OtherLead')
        other.profile.project = Project.objects.create(name='Other project')
        other.profile.save()
        client = Client()
        client.force_login(other)
        self.assertEqual(self.review_decision(c, 'approve', client=client).status_code, 404)
        for selection in [[], ['unknown']]:
            self.review_decision(c, 'partial', selection)
            c.refresh_from_db()
            self.assertEqual(c.section_approvals, {})
        self.review_decision(c, 'partial', ['general'])
        self.review_decision(c, 'partial', ['general'])
        c.refresh_from_db()
        self.assertEqual(len(c.review_history), 2)

    def test_selected_section_required_fields_gate_approval(self):
        c = self.review_fixture()
        # Submitted engineers are read-only; seed incomplete synthetic evidence to test the review gate.
        self.assertEqual(self.post_answers(c, {'volts': ''}).status_code, 400)
        c.answer_data['custom']['volts'] = ''
        c.save(update_fields=['answer_data', 'updated_at'])
        self.review_decision(c, 'approve')
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.REVIEW)
        self.assertEqual(c.section_approvals,{})
        from .quality_workflow import review_quality
        review_quality(self.quality,c,'approve',expected_updated_at=c.updated_at.isoformat())
        self.review_decision(c, 'partial', ['general'])
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.PARTIAL)
        self.review_decision(c, 'partial', ['power'])
        c.refresh_from_db()
        self.assertNotIn('power', c.section_approvals)

    def test_request_changes_preserves_locks_and_legacy_review_cannot_bypass(self):
        c = self.review_fixture()
        self.review_decision(c, 'partial', ['general'])
        self.review_decision(c, 'changes', comment='Recheck power voltage')
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.DRAFT)
        self.assertIn('general', c.section_approvals)
        self.assertEqual(self.post_answers(c, {'height': '60'}).status_code, 400)
        self.l.post(reverse('checklist_review_update', args=[c.pk]), {'status': 'FINAL', 'comment': 'Attempt to bypass submission'})
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.DRAFT)
        self.e.post(reverse('custom_checklist_submit', args=[c.pk]))
        c.refresh_from_db()
        self.assertEqual(c.status, Checklist.Status.PARTIAL)


    def test_approval_history_in_latest_excel_export(self):
        c = self.review_fixture()
        self.review_decision(c, 'partial', ['general'], comment='Foundation verified')
        response = self.e.get(reverse('custom_checklist_download', args=[c.pk]))
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook['Section approvals']['B2'].value, 'Approved')
        self.assertEqual(workbook['Section approvals']['B3'].value, 'Pending')
        self.assertEqual(workbook['Review history']['C3'].value, 'partial')
        self.assertEqual(workbook['Review history']['E3'].value, 'Foundation verified')
